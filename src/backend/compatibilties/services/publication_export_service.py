import logging
import os
import re
import time
from collections.abc import Callable
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from config import settings
from services.job_store import JobStore
from services.supabase_publications_store import supabase_publications_store

logger = logging.getLogger(__name__)

_ILLEGAL_EXCEL_CHARACTERS = re.compile(
    r"[\x00-\x08\x0B\x0C\x0E-\x1F]"
)
_MAX_EXCEL_CELL_LENGTH = 32_767
_EXPORT_ARTIFACT_NAME = re.compile(
    r"^[0-9a-fA-F-]{36}_publicaciones_\d{4}-\d{2}-\d{2}"
    r"(?:\.xlsx(?:\.tmp)?|\.checkpoint\.jsonl)$"
)


def _excel_text(value: Any) -> str:
    text = _ILLEGAL_EXCEL_CHARACTERS.sub("", str(value or ""))
    return text[:_MAX_EXCEL_CELL_LENGTH]


class PublicationExportService:
    async def export(
        self,
        *,
        job_id: str,
        user_id: str,
        creation_date: str,
        heartbeat: Callable[[], None] | None = None,
    ) -> dict[str, Any]:
        self._cleanup_expired_artifacts()
        JobStore.update(
            job_id,
            status="processing",
            progress=10,
            processed_rows=0,
            last_error=None,
            message="Leyendo publicaciones desde la base de datos...",
        )

        publications = await supabase_publications_store.list_by_creation_date(
            creation_date,
            seller_id=user_id,
        )
        total = len(publications)
        if total == 0:
            raise ValueError(
                f"No existen publicaciones con fecha de creacion {creation_date}."
            )

        if heartbeat is not None:
            heartbeat()
        JobStore.update(
            job_id,
            status="processing",
            progress=70,
            total_rows=total,
            message=(
                f"Construyendo Excel con {total} publicaciones desde la "
                "base de datos..."
            ),
        )

        output_path = self._output_path(
            job_id=job_id,
            creation_date=creation_date,
        )
        self._build_excel_atomically(
            publications=publications,
            output_path=output_path,
        )

        if heartbeat is not None:
            heartbeat()

        filename = f"publicaciones_{creation_date}.xlsx"
        summary = {
            "total_rows": total,
            "source": "database",
            "api_items_queried": 0,
        }
        JobStore.update(
            job_id,
            status="success",
            progress=100,
            processed_rows=total,
            result_path=output_path,
            output_filename=filename,
            summary=summary,
            last_error=None,
            message=(
                f"Excel listo: {total} publicaciones exportadas desde la "
                "base de datos."
            ),
        )
        return summary

    def _build_excel_atomically(
        self,
        *,
        publications: list[dict[str, Any]],
        output_path: str,
    ) -> None:
        workbook, worksheet = self._create_workbook()
        for publication in publications:
            worksheet.append(
                [
                    _excel_text(publication.get("mlc")),
                    _excel_text(publication.get("sku")),
                    _excel_text(publication.get("titulo")),
                ]
            )

        os.makedirs(settings.upload_dir, exist_ok=True)
        worksheet.auto_filter.ref = f"A1:C{len(publications) + 1}"
        temporary_output_path = f"{output_path}.tmp"
        try:
            workbook.save(temporary_output_path)
            os.replace(temporary_output_path, output_path)
        finally:
            try:
                os.remove(temporary_output_path)
            except FileNotFoundError:
                pass

    @staticmethod
    def _cleanup_expired_artifacts() -> None:
        uploads_root = Path(settings.upload_dir)
        if not uploads_root.exists():
            return

        cutoff = (
            time.time()
            - int(settings.ml_publication_export_artifact_ttl_seconds)
        )
        for path in uploads_root.iterdir():
            if (
                not path.is_file()
                or not _EXPORT_ARTIFACT_NAME.fullmatch(path.name)
            ):
                continue
            try:
                if path.stat().st_mtime < cutoff:
                    path.unlink()
            except OSError:
                logger.exception(
                    "[PUBLICATION_EXPORT][ARTIFACT_CLEANUP_ERROR] "
                    "path=%s",
                    path,
                )

    @staticmethod
    def _create_workbook():
        workbook = Workbook(write_only=True)
        worksheet = workbook.create_sheet("Publicaciones")
        worksheet.freeze_panes = "A2"
        worksheet.column_dimensions["A"].width = 20
        worksheet.column_dimensions["B"].width = 24
        worksheet.column_dimensions["C"].width = 55

        header_fill = PatternFill("solid", fgColor="3483FA")
        header_font = Font(color="FFFFFF", bold=True)
        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        header = []
        for value in ("MLC", "SKU", "TITULO"):
            cell = WriteOnlyCell(worksheet, value=value)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            header.append(cell)
        worksheet.append(header)

        return workbook, worksheet

    @staticmethod
    def _output_path(*, job_id: str, creation_date: str) -> str:
        return os.path.join(
            settings.upload_dir,
            f"{job_id}_publicaciones_{creation_date}.xlsx",
        )


publication_export_service = PublicationExportService()
