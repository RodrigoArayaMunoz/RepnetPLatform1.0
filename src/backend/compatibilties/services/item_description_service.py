import asyncio
import json
import logging
import os
import re
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from config import settings
from services.compatibility_service import JobMetrics, WRITE_RATE_LIMITER
from services.excel_service import normalize_for_compare
from services.job_store import JobStore
from services.ml_client import ml_client
from services.process_chunking_service import chunk_sequence, count_chunks
from services.redis_rate_limiter import CombinedRateLimiter, RedisWindowRateLimiter

logger = logging.getLogger(__name__)
MLC_COLUMN = "MLC"
DESCRIPTION_COLUMN = "DESCRIPCION A"

ITEM_DESCRIPTION_RATE_LIMITER = RedisWindowRateLimiter(
    redis_url=settings.redis_url,
    namespace="ml:item_description",
    requests_per_second=settings.ml_item_description_requests_per_second,
    max_requests_per_window=settings.ml_item_description_max_requests_per_window,
    window_seconds=settings.ml_item_description_window_seconds,
)
ITEM_DESCRIPTION_WRITE_RATE_LIMITER = CombinedRateLimiter(
    ITEM_DESCRIPTION_RATE_LIMITER, WRITE_RATE_LIMITER,
)


def load_item_description_rows(file_path: str) -> list[dict[str, Any]]:
    if Path(file_path).suffix.lower() != ".xlsx":
        raise ValueError("Las descripciones por MLC requieren un Excel .xlsx con Hoja1")
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if "Hoja1" not in workbook.sheetnames:
            raise ValueError("El archivo de descripciones debe contener Hoja1")
        values = workbook["Hoja1"].iter_rows(values_only=True)
        headers = list(next(values, ()))
        normalized = [normalize_for_compare(str(value or "").strip()) for value in headers]
        indexes = []
        for label in (MLC_COLUMN, DESCRIPTION_COLUMN):
            matches = [i for i, header in enumerate(normalized) if header == normalize_for_compare(label)]
            if len(matches) != 1:
                raise ValueError(f"Hoja1 debe tener exactamente una columna {label}")
            indexes.append(matches[0])
        rows = []
        for excel_row, cells in enumerate(values, start=2):
            raw_id, text = (cells[i] if i < len(cells) else None for i in indexes)
            if raw_id is None and text is None:
                continue
            item_id = str(raw_id or "").strip().upper()
            # Keep valid text byte-for-byte; invalid cell types still need a
            # serializable value in the per-row error report.
            report_text = text if isinstance(text, str) else ("" if text is None else str(text))
            row = {"item_id": item_id, "plain_text": report_text,
                   "original_row_index": excel_row - 2, "excel_row": excel_row}
            if not re.fullmatch(r"MLC\d+", item_id):
                row["validation_error"] = "MLC inválido: debe tener el formato MLC seguido de números"
            elif not isinstance(text, str) or not text.strip():
                row["validation_error"] = "DESCRIPCION A debe contener texto y no puede estar vacía"
            rows.append(row)
    finally:
        workbook.close()
    if not rows:
        raise ValueError("Hoja1 no contiene filas para procesar")

    groups = defaultdict(list)
    for index, row in enumerate(rows):
        if not row.get("validation_error"):
            groups[row["item_id"]].append(index)
    for item_id, indexes in groups.items():
        if len({rows[index]["plain_text"] for index in indexes}) > 1:
            for index in indexes:
                rows[index]["validation_error"] = f"MLC duplicado con descripciones diferentes: {item_id}"
        else:
            for index in indexes[1:]:
                rows[index]["duplicate_of"] = indexes[0]
    return rows


def _same_text(first: Any, second: str) -> bool:
    def normalized(text: str) -> str:
        # Live GET responses omit terminal newlines from the submitted text.
        # Ignore that normalization only when comparing; preserve the payload,
        # leading spaces, internal blank lines and other trailing characters.
        return text.replace("\r\n", "\n").replace("\r", "\n").rstrip("\n")
    return isinstance(first, str) and normalized(first) == normalized(second)


async def _process_item_description_row(row: dict, user_id: str, metrics: JobMetrics) -> dict:
    result = {**row, "ok": False}
    if row.get("validation_error"):
        return {**result, "reason": row["validation_error"], "error_code": "INVALID_ROW"}
    item_id = row["item_id"]
    text = row["plain_text"]
    try:
        try:
            current = await ml_client.get_item_description(
                None, item_id, user_id=user_id,
                rate_limiter=ITEM_DESCRIPTION_RATE_LIMITER, metrics=metrics,
            )
            exists = True
        except HTTPException as exc:
            if exc.status_code != 404:
                raise
            current = {}
            exists = False

        if exists and _same_text(current.get("plain_text"), text):
            return {**result, "ok": True, "action": "unchanged", "reason": "La descripción ya coincide con el Excel"}

        try:
            response = await ml_client.write_item_description(
                item_id, text, exists=exists, user_id=user_id,
                rate_limiter=ITEM_DESCRIPTION_WRITE_RATE_LIMITER, metrics=metrics,
            )
        except HTTPException as write_error:
            # A POST may have been applied before a timeout, so its retry can
            # return 400 (description already exists). Verify instead of
            # reporting a false failure or blindly replaying the POST.
            if write_error.status_code not in {400, 502}:
                raise
            try:
                response = await ml_client.get_item_description(
                    None, item_id, user_id=user_id,
                    rate_limiter=ITEM_DESCRIPTION_RATE_LIMITER, metrics=metrics,
                )
            except HTTPException:
                raise write_error
            if not _same_text(response.get("plain_text"), text):
                raise write_error
        # Some write responses contain only metadata. Read back in that case;
        # a 2xx without confirmation must not be reported as a successful change.
        if "plain_text" not in response:
            response = await ml_client.get_item_description(
                None, item_id, user_id=user_id,
                rate_limiter=ITEM_DESCRIPTION_RATE_LIMITER, metrics=metrics,
            )
        if not _same_text(response.get("plain_text"), text):
            return {**result, "reason": "Mercado Libre no confirmó el texto enviado",
                    "error_code": "DESCRIPTION_NOT_CONFIRMED"}
        return {**result, "ok": True, "action": "updated" if exists else "created",
                "reason": "Descripción actualizada" if exists else "Descripción creada"}
    except HTTPException as exc:
        return {**result, "reason": str(exc.detail), "error_code": f"HTTP_{exc.status_code}"}
    except Exception as exc:
        logger.exception("[ITEM_DESCRIPTION][ERROR] item_id=%s", item_id)
        return {**result, "reason": str(exc), "error_code": "UNEXPECTED_EXCEPTION"}


async def process_item_description_job(job_id: str, user_id: str, file_path: str) -> dict:
    rows = load_item_description_rows(file_path)
    metrics = JobMetrics()
    chunk_size = settings.item_description_chunk_size
    concurrency = settings.item_description_max_concurrency
    progress_every = settings.job_progress_update_every
    duplicates = defaultdict(list)
    entries = []
    for index, row in enumerate(rows):
        if "duplicate_of" in row:
            duplicates[row["duplicate_of"]].append(index)
        else:
            entries.append((index, row))
    total_chunks = count_chunks(len(entries), chunk_size)
    results = [None] * len(rows)
    completed = 0
    completed_unique = 0
    logger.info("[ITEM_DESCRIPTION][POLICY] rows=%s chunk_size=%s concurrency=%s endpoint_rps=%s endpoint_window=%s/%ss",
                len(rows), chunk_size, concurrency, ITEM_DESCRIPTION_RATE_LIMITER.requests_per_second,
                ITEM_DESCRIPTION_RATE_LIMITER.max_requests_per_window, ITEM_DESCRIPTION_RATE_LIMITER.window_seconds)
    JobStore.update(job_id, status="processing", progress=5, total_rows=len(rows),
                    processed_rows=0, processed_unique_rows=0,
                    total_unique_rows=len(entries), total_chunks=total_chunks, completed_chunks=0,
                    message="Preparando descripciones por MLC...")

    async def worker(index: int, row: dict, semaphore: asyncio.Semaphore) -> None:
        nonlocal completed, completed_unique
        async with semaphore:
            result = await _process_item_description_row(row, user_id, metrics)
            results[index] = result
            completed_unique += 1
            resolved_indexes = [index, *duplicates[index]]
            for duplicate_index in duplicates[index]:
                metrics.cache_hits += 1
                results[duplicate_index] = {**result, **rows[duplicate_index], "action": "duplicate"}
            for _ in resolved_indexes:
                completed += 1
                if completed % progress_every == 0 or completed == len(rows):
                    JobStore.update(job_id, processed_rows=completed, processed_unique_rows=completed_unique,
                                    progress=min(95, 10 + int(completed / len(rows) * 85)),
                                    message=f"Descripciones: {completed}/{len(rows)} filas procesadas")

    semaphore = asyncio.Semaphore(concurrency)
    for chunk_number, (_, chunk) in enumerate(chunk_sequence(entries, chunk_size), start=1):
        JobStore.update(job_id, message=f"Descripciones: bloque {chunk_number}/{total_chunks}")
        await asyncio.gather(*(worker(index, row, semaphore) for index, row in chunk))
        JobStore.update(job_id, completed_chunks=chunk_number)

    success_count = sum(bool(result["ok"]) for result in results)
    summary = {
        "processed_rows": len(rows), "unique_rows": len(entries),
        "success_count": success_count, "error_count": len(rows) - success_count,
        "created_items": sum(result.get("action") == "created" for result in results),
        "updated_items": sum(result.get("action") == "updated" for result in results),
        "unchanged_items": sum(result.get("action") == "unchanged" for result in results),
        "duplicate_rows": sum("duplicate_of" in row for row in rows),
        "failed_item_ids": list(dict.fromkeys(result["item_id"] for result in results if not result["ok"] and result["item_id"])),
        "exported_rows_total": len(rows), "metrics": metrics.to_dict(),
    }
    os.makedirs(settings.upload_dir, exist_ok=True)
    result_path = os.path.join(settings.upload_dir, f"{job_id}_item_descriptions_result.json")
    with open(f"{result_path}.tmp", "w", encoding="utf-8") as output:
        json.dump(results, output, ensure_ascii=False)
    os.replace(f"{result_path}.tmp", result_path)
    JobStore.update(job_id, status="success", progress=100, processed_rows=len(rows),
                    processed_unique_rows=len(entries),
                    result_path=result_path, summary=summary,
                    message="Procesamiento de descripciones finalizado")
    return {"results": results, "summary": summary}


def build_item_description_excel(results: list[dict]) -> BytesIO:
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Hoja1")
    worksheet.freeze_panes = "A2"
    for column, width in (("A", 24), ("B", 75), ("C", 24), ("D", 60), ("E", 15)):
        worksheet.column_dimensions[column].width = width
    headers = [MLC_COLUMN, DESCRIPTION_COLUMN, "RESULTADO", "DETALLE", "FILA EXCEL"]
    cells = [WriteOnlyCell(worksheet, value=header) for header in headers]
    for cell in cells:
        cell.font = Font(bold=True)
    worksheet.append(cells)
    actions = {"created": "Creada", "updated": "Actualizada", "unchanged": "Sin cambios", "duplicate": "Duplicada"}
    for result in results:
        values = [result.get("item_id", ""), result.get("plain_text") or "",
                  actions.get(result.get("action"), "Procesada") if result.get("ok") else "Error",
                  result.get("reason", ""), str(result.get("excel_row", ""))]
        cells = []
        for value in values:
            cell = WriteOnlyCell(worksheet, value=str(value))
            cell.data_type = "s"
            cells.append(cell)
        worksheet.append(cells)
    worksheet.auto_filter.ref = f"A1:E{len(results) + 1}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
