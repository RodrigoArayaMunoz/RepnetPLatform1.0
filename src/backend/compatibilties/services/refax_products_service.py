import asyncio
import json
import logging
import re
from collections.abc import Callable
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import settings
from services.refax_connection_service import (
    RefaxConnectionError,
    RefaxConnectionService,
    refax_connection_service,
    refax_test_connection_service,
)

logger = logging.getLogger(__name__)
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ProgressCallback = Callable[[int, str, str], None]
PRODUCT_COLUMNS = (
    ("SKU", "numero_refax", 25),
    ("PRECIO", "precio", 18),
    ("STOCK", "stock", 20),
    ("NOMBRE_PRODUCTO", "nombre_producto", 35),
    ("GLOSA_1", "glosa_1", 45),
    ("GLOSA_2", "glosa_2", 45),
    ("GLOSA_3", "glosa_3", 45),
    ("FACTOR", "factor", 15),
    ("MARCA_PRODUCTO", "marca_producto", 25),
    ("ORIGEN", "origen", 20),
    ("RUBRO_COMERCIAL", "rubro_comercial", 25),
    ("IMAGENES_URL", "imagenes_url", 80),
    ("CODIGO_OEM", "codigo_oem", 25),
    ("MARCA_OEM", "marca_oem", 25),
    ("CODIGO_FABRICA", "codigo_fabrica", 25),
    ("FABRICA", "fabrica", 25),
    ("APLICACIONES", "aplicaciones", 80),
    ("PRECIO_OFERTA_WEB", "precio_oferta_web", 25),
    ("CANTIDAD_OFERTA_WEB", "cantidad_oferta_web", 25),
)
_NUMERIC_COLUMNS = {"PRECIO", "FACTOR", "PRECIO_OFERTA_WEB", "CANTIDAD_OFERTA_WEB"}
_FIELD_ALIASES = {
    "aplicaciones": ("aplicaciones", "aplicacion"),
    "imagenes_url": ("imagenes_url", "imagen_url"),
}
_ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


class RefaxProductsError(RuntimeError):
    pass


@dataclass(frozen=True)
class RefaxProductsDownload:
    content: bytes
    content_type: str


class RefaxProductsService:
    PRODUCTS_PATH = "/api/Productos/Listado"
    MAX_ATTEMPTS = 3
    RETRYABLE_STATUSES = {408, 429, 500, 502, 503, 504}

    def __init__(
        self,
        connection_service: RefaxConnectionService | None = None,
        *,
        environment: str = "production",
    ) -> None:
        if environment not in {"production", "test"}:
            raise ValueError("Ambiente REFAX no valido")
        self.environment = environment
        self.connection_service = connection_service if connection_service is not None else (
            refax_test_connection_service
            if environment == "test"
            else refax_connection_service
        )

    async def _request_products(
        self, token: str, on_progress: ProgressCallback | None = None
    ) -> httpx.Response:
        prefix = "refax_test" if self.environment == "test" else "refax"
        base_url = getattr(settings, f"{prefix}_api_base_url")
        provider_code = getattr(settings, f"{prefix}_provider_code")
        if not provider_code:
            raise RefaxProductsError(f"Configura {prefix.upper()}_PROVIDER_CODE en el backend")
        url = f"{base_url.rstrip('/')}{self.PRODUCTS_PATH}"
        timeout = httpx.Timeout(
            settings.refax_products_http_timeout_seconds,
            connect=min(15.0, settings.refax_products_http_timeout_seconds),
        )
        for attempt in range(1, self.MAX_ATTEMPTS + 1):
            delay = float(2 ** (attempt - 1))
            try:
                async with httpx.AsyncClient(timeout=timeout) as client:
                    request_options = {
                        "params": {"codigo": provider_code},
                        "headers": {
                            "Authorization": f"Bearer {token}",
                            "Accept": "application/json",
                        },
                    }
                    if on_progress is None:
                        response = await client.get(url, **request_options)
                    else:
                        on_progress(5, "requesting", f"Consultando productos en REFAX (intento {attempt})…")
                        request_options["headers"]["Accept-Encoding"] = "identity"
                        async with client.stream("GET", url, **request_options) as upstream:
                            if upstream.status_code != 200:
                                await upstream.aread()
                                response = upstream
                            else:
                                length = upstream.headers.get("content-length", "")
                                total = int(length) if length.isdigit() and int(length) > 0 else None
                                chunks = []
                                received = 0
                                on_progress(10, "receiving", "Recibiendo productos de REFAX…")
                                async for chunk in upstream.aiter_bytes(chunk_size=64 * 1024):
                                    chunks.append(chunk)
                                    received += len(chunk)
                                    percentage = 10 + min(45, int(45 * received / total)) if total else 10
                                    on_progress(
                                        percentage, "receiving",
                                        f"Recibidos {received / (1024 * 1024):.1f} MB de productos desde REFAX…",
                                    )
                                response = httpx.Response(
                                    upstream.status_code, content=b"".join(chunks),
                                    headers={"Content-Type": upstream.headers.get("content-type", "application/json")},
                                )
            except httpx.HTTPError as error:
                # No registrar cuerpos, URL con codigo ni encabezados con token.
                logger.warning(
                    "[REFAX_PRODUCTS] attempt=%s/%s error_type=%s",
                    attempt, self.MAX_ATTEMPTS, type(error).__name__,
                )
                if attempt == self.MAX_ATTEMPTS:
                    message = (
                        "REFAX excedió el tiempo de espera al entregar los productos. "
                        if isinstance(error, httpx.TimeoutException)
                        else "Se interrumpió la comunicación con REFAX. "
                    )
                    raise RefaxProductsError(
                        message + "Se realizaron 3 intentos; vuelve a intentar la descarga."
                    ) from error
            else:
                if (
                    response.status_code not in self.RETRYABLE_STATUSES
                    or attempt == self.MAX_ATTEMPTS
                ):
                    return response
                retry_after = response.headers.get("Retry-After")
                if retry_after:
                    # Si el proveedor exige una espera larga o una fecha, dejar
                    # el reintento al usuario sin enviar peticiones prematuras.
                    if not retry_after.isdigit() or int(retry_after) > 30:
                        return response
                    delay = max(delay, float(retry_after))
                logger.warning(
                    "[REFAX_PRODUCTS] attempt=%s/%s http_status=%s",
                    attempt, self.MAX_ATTEMPTS, response.status_code,
                )
            if on_progress:
                on_progress(5, "retrying", "Reintentando la comunicación con REFAX…")
            await asyncio.sleep(delay)

        raise RefaxProductsError("No se pudo completar la consulta a REFAX")

    @staticmethod
    def _number(value, *, row_number: int, column: str):
        if isinstance(value, bool) or value is None:
            raise RefaxProductsError(
                f"REFAX entregó {column} inválido en el producto {row_number}."
            )
        try:
            number = Decimal(str(value).strip())
        except InvalidOperation as error:
            raise RefaxProductsError(
                f"REFAX entregó {column} inválido en el producto {row_number}."
            ) from error
        if not number.is_finite():
            raise RefaxProductsError(
                f"REFAX entregó {column} inválido en el producto {row_number}."
            )
        return number

    @staticmethod
    def _decimal_text(number: Decimal) -> str:
        text = format(number, "f")
        return text.rstrip("0").rstrip(".") if "." in text else text

    @classmethod
    def _field_text(
        cls, value, *, row_number: int, column: str, required: bool = False,
    ) -> str:
        if isinstance(value, list):
            # Cualquier campo puede tener varios valores. Conservar su orden
            # y la estructura JSON de objetos y listas anidados.
            text = "\n".join(
                json.dumps(item, ensure_ascii=False)
                if isinstance(item, (dict, list))
                else cls._field_text(item, row_number=row_number, column=column)
                for item in value
            )
            if required and not text.strip():
                raise RefaxProductsError(
                    f"REFAX entregó {column} inválido en el producto {row_number}."
                )
            return text
        if isinstance(value, dict):
            return json.dumps(value, ensure_ascii=False)
        if value is None or value == "":
            if required:
                raise RefaxProductsError(
                    f"REFAX entregó {column} inválido en el producto {row_number}."
                )
            return ""
        if column == "SKU":
            if isinstance(value, (str, int)) and not isinstance(value, bool) and str(value).strip():
                return str(value)
            raise RefaxProductsError(
                f"REFAX entregó SKU inválido en el producto {row_number}."
            )
        if column in _NUMERIC_COLUMNS or (column == "STOCK" and not isinstance(value, str)):
            return cls._decimal_text(
                cls._number(value, row_number=row_number, column=column)
            )
        if isinstance(value, bool):
            return json.dumps(value)
        if isinstance(value, (str, int, float)):
            return str(value)
        raise RefaxProductsError(
            f"REFAX entregó {column} inválido en el producto {row_number}."
        )

    @staticmethod
    def _product_value(product: dict, field: str):
        # PRODUCCION usa nombres plurales; TEST devuelve estos campos en singular.
        for name in _FIELD_ALIASES.get(field, (field,)):
            value = product.get(name)
            if value not in (None, "", []):
                return value
        return None

    @staticmethod
    def _validate_excel_text(value: str, *, row_number: int, column: str) -> None:
        # Excel recorta silenciosamente celdas largas. No entregar datos incompletos.
        if len(value) > 32767:
            raise RefaxProductsError(
                f"{column} del producto {row_number} supera el límite de "
                "32.767 caracteres por celda de Excel."
            )
        if _ILLEGAL_EXCEL_CHARACTERS.search(value):
            raise RefaxProductsError(
                f"REFAX entregó caracteres incompatibles con Excel en "
                f"{column} del producto {row_number}."
            )

    @staticmethod
    def _mark_export_as_intentional_text(content: bytes, last_row: int) -> bytes:
        # openpyxl no serializa ignoredErrors. Añadir la regla OOXML solo
        # a las celdas exportadas, sin modificar las preferencias de Excel.
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with BytesIO(content) as source, BytesIO() as output:
            with ZipFile(source) as original, ZipFile(output, "w") as updated:
                for entry in original.infolist():
                    data = original.read(entry.filename)
                    if entry.filename == "xl/worksheets/sheet1.xml":
                        sheet = ElementTree.fromstring(data)
                        ignored = ElementTree.SubElement(sheet, f"{{{namespace}}}ignoredErrors")
                        ElementTree.SubElement(
                            ignored,
                            f"{{{namespace}}}ignoredError",
                            sqref=f"A2:{get_column_letter(len(PRODUCT_COLUMNS))}{last_row}",
                            numberStoredAsText="1",
                        )
                        data = ElementTree.tostring(sheet, encoding="utf-8")
                    updated.writestr(entry, data)
            return output.getvalue()

    @classmethod
    def _build_excel(
        cls, content: bytes, on_progress: ProgressCallback | None = None
    ) -> bytes:
        # Se ejecuta fuera del event loop para mantener disponible el backend.
        try:
            products = json.loads(content)
        except (ValueError, UnicodeDecodeError) as error:
            raise RefaxProductsError(
                "REFAX entregó una respuesta inválida; vuelve a intentar la descarga."
            ) from error
        if not isinstance(products, list):
            raise RefaxProductsError("REFAX no entregó un listado de productos válido.")
        if not products:
            raise RefaxProductsError("REFAX entregó un listado de productos vacío.")

        total = len(products)
        if on_progress:
            on_progress(60, "building", f"Generando Excel: 0 de {total:,} productos.")

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Productos"
        worksheet.append([header for header, _field, _width in PRODUCT_COLUMNS])
        worksheet.freeze_panes = "A2"
        for column_index, (_header, _field, width) in enumerate(PRODUCT_COLUMNS, start=1):
            column = get_column_letter(column_index)
            worksheet.column_dimensions[column].width = width
            worksheet.column_dimensions[column].number_format = "@"
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="6B8CFF")
        data_alignment = Alignment(wrap_text=True, vertical="top")

        try:
            for index, product in enumerate(products, start=1):
                if not isinstance(product, dict) or any(
                    key not in product for key in ("numero_refax", "precio", "stock")
                ):
                    raise RefaxProductsError(
                        f"Faltan SKU, PRECIO o STOCK en el producto {index} de REFAX."
                    )
                values = [
                    cls._field_text(
                        cls._product_value(product, field), row_number=index, column=header,
                        required=header in {"SKU", "PRECIO"},
                    )
                    for header, field, _width in PRODUCT_COLUMNS
                ]
                for (header, _field, _width), value in zip(PRODUCT_COLUMNS, values):
                    cls._validate_excel_text(value, row_number=index, column=header)
                worksheet.append(values)
                for column_index in range(1, len(PRODUCT_COLUMNS) + 1):
                    cell = worksheet.cell(row=index + 1, column=column_index)
                    # Texto real y formato Texto; evitar fórmulas y preservar SKU.
                    cell.data_type = "s"
                    cell.number_format = "@"
                    cell.alignment = data_alignment
                if on_progress and (index % max(1, total // 30) == 0 or index == total):
                    on_progress(
                        60 + int(30 * index / total), "building",
                        f"Generando Excel: {index:,} de {total:,} productos.",
                    )
            last_column = get_column_letter(len(PRODUCT_COLUMNS))
            worksheet.auto_filter.ref = f"A1:{last_column}{len(products) + 1}"
            with BytesIO() as output:
                if on_progress:
                    on_progress(90, "packaging", "Guardando las hojas del Excel…")
                workbook.save(output)
                if on_progress:
                    on_progress(93, "formatting", "Aplicando el formato final del archivo…")
                result = cls._mark_export_as_intentional_text(
                    output.getvalue(), last_row=len(products) + 1
                )
                if on_progress:
                    on_progress(95, "ready", "Excel preparado. Enviando el archivo…")
                return result
        finally:
            workbook.close()

    async def download(
        self, on_progress: ProgressCallback | None = None
    ) -> RefaxProductsDownload:
        if on_progress:
            on_progress(0, "connecting", "Verificando la conexión con REFAX…")
        try:
            token = await self.connection_service.get_valid_token()
        except RefaxConnectionError as error:
            raise RefaxProductsError(str(error)) from error

        if on_progress:
            on_progress(5, "requesting", "Conexión verificada. Consultando productos…")
        response = await self._request_products(token, on_progress) if on_progress else await self._request_products(token)

        # Un token puede ser revocado por REFAX antes de su expiracion declarada.
        # En ese caso se obtiene uno nuevo y se reintenta una unica vez.
        if response.status_code in (401, 403):
            try:
                await self.connection_service.connect()
                token = await self.connection_service.get_valid_token()
            except RefaxConnectionError as error:
                raise RefaxProductsError(str(error)) from error
            response = await self._request_products(token, on_progress) if on_progress else await self._request_products(token)

        if response.status_code >= 400:
            raise RefaxProductsError(
                f"REFAX no pudo entregar los productos (HTTP {response.status_code})"
            )
        if not response.content:
            raise RefaxProductsError("REFAX entrego un listado de productos vacio")

        if on_progress:
            on_progress(58, "validating", "Listado recibido. Validando los productos…")

        return RefaxProductsDownload(
            content=await asyncio.to_thread(self._build_excel, response.content, on_progress),
            content_type=XLSX_CONTENT_TYPE,
        )


refax_products_service = RefaxProductsService()
refax_test_products_service = RefaxProductsService(environment="test")
