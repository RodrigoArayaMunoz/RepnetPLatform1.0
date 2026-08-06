import os
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

import httpx
from fastapi import HTTPException, Request
from openpyxl import load_workbook

from celery_app import celery_app
from config import settings
from routers import publications_router
from services.job_store import JobStore
from services.ml_client import MercadoLibreClient, ml_client
from services.publication_export_service import PublicationExportService
from services.publication_export_store import PublicationExportStore
from services.supabase_publications_store import supabase_publications_store


class _FakeRedisClient:
    async def aclose(self):
        return None


class _FakeLimiter:
    def __init__(self, *args, **kwargs):
        self.client = _FakeRedisClient()
        self.acquire_count = 0
        self.penalties = []

    async def acquire(self):
        self.acquire_count += 1

    async def penalize(self, cooldown_seconds):
        self.penalties.append(cooldown_seconds)


class MercadoLibreClientRateLimitTests(unittest.IsolatedAsyncioTestCase):
    async def test_429_pauses_shared_limiter_and_respects_retry_after(self):
        client = MercadoLibreClient()
        client.client = Mock()
        client.client.request = AsyncMock(
            side_effect=[
                httpx.Response(
                    429,
                    headers={"Retry-After": "120"},
                    json={"message": "too_many_requests"},
                    request=httpx.Request(
                        "GET",
                        "https://api.mercadolibre.com/items/MLC1",
                    ),
                ),
                httpx.Response(
                    200,
                    json={"id": "MLC1", "title": "Titulo"},
                    request=httpx.Request(
                        "GET",
                        "https://api.mercadolibre.com/items/MLC1",
                    ),
                ),
            ]
        )
        limiter = _FakeLimiter()

        with (
            patch.object(
                client,
                "get_valid_token",
                new=AsyncMock(return_value="token"),
            ),
            patch.object(settings, "ml_retry_attempts", 2),
            patch(
                "services.ml_client.asyncio.sleep",
                new=AsyncMock(),
            ) as sleep,
            patch("services.ml_client.random.uniform", return_value=0),
        ):
            payload = await client.request(
                "GET",
                "/items/MLC1",
                user_id="2682261950",
                rate_limiter=limiter,
            )

        self.assertEqual(payload["title"], "Titulo")
        self.assertEqual(limiter.acquire_count, 2)
        self.assertEqual(limiter.penalties, [120.0])
        sleep.assert_awaited_once_with(120.0)


class PublicationExcelTests(unittest.IsolatedAsyncioTestCase):
    async def test_export_writes_only_required_columns_from_database(self):
        service = PublicationExportService()
        publications = [
            {"mlc": "MLC123", "sku": "SKU-1", "titulo": "Titulo 1"},
            {"mlc": "MLC456", "sku": "SKU-2", "titulo": "Titulo 2"},
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            with (
                patch.object(settings, "upload_dir", temp_dir),
                patch.object(
                    supabase_publications_store,
                    "list_by_creation_date",
                    new=AsyncMock(return_value=publications),
                ),
                patch.object(
                    supabase_publications_store,
                    "upsert_rows",
                    new=AsyncMock(),
                ) as upsert_rows,
                patch.object(
                    ml_client,
                    "request",
                    new=AsyncMock(),
                ) as ml_request,
                patch.object(JobStore, "update", new=Mock()),
            ):
                summary = await service.export(
                    job_id="job-1",
                    user_id="2682261950",
                    creation_date="2026-07-24",
                )

            output_path = os.path.join(
                temp_dir,
                "job-1_publicaciones_2026-07-24.xlsx",
            )
            workbook = load_workbook(output_path, read_only=True)
            rows = list(
                workbook["Publicaciones"].iter_rows(values_only=True)
            )
            workbook.close()

        self.assertEqual(rows[0], ("MLC", "SKU", "TITULO"))
        self.assertEqual(rows[1], ("MLC123", "SKU-1", "Titulo 1"))
        self.assertEqual(rows[2], ("MLC456", "SKU-2", "Titulo 2"))
        self.assertEqual(summary["total_rows"], 2)
        self.assertEqual(summary["source"], "database")
        self.assertEqual(summary["api_items_queried"], 0)
        ml_request.assert_not_awaited()
        upsert_rows.assert_not_awaited()

    async def test_export_rejects_empty_database_result(self):
        service = PublicationExportService()

        with tempfile.TemporaryDirectory() as temp_dir:
            with (
                patch.object(settings, "upload_dir", temp_dir),
                patch.object(
                    supabase_publications_store,
                    "list_by_creation_date",
                    new=AsyncMock(return_value=[]),
                ),
                patch.object(JobStore, "update", new=Mock()),
            ):
                with self.assertRaises(ValueError):
                    await service.export(
                        job_id="job-empty",
                        user_id="2682261950",
                        creation_date="2026-07-24",
                    )

    def test_export_task_uses_dedicated_export_queue(self):
        routes = celery_app.conf.task_routes
        self.assertEqual(
            routes["tasks.export_publications_job"]["queue"],
            "publication_exports",
        )

    def test_export_worker_accepts_four_jobs_in_parallel_by_default(self):
        self.assertEqual(
            settings.publication_export_worker_concurrency,
            4,
        )


class PublicationExportOwnershipTests(unittest.TestCase):
    @staticmethod
    def _request(user_id: str) -> Request:
        request = Mock(spec=Request)
        request.state = SimpleNamespace(
            supabase_user={"id": user_id}
        )
        return request

    def test_reference_is_isolated_by_authenticated_user(self):
        first = PublicationExportStore._reference_key(
            "supabase-user-a",
            "2682261950",
            "2026-07-24",
        )
        second = PublicationExportStore._reference_key(
            "supabase-user-b",
            "2682261950",
            "2026-07-24",
        )
        self.assertNotEqual(first, second)

    def test_user_cannot_read_another_users_export(self):
        job = {
            "id": "job-a",
            "job_type": "publication_export",
            "requested_by_user_id": "supabase-user-a",
        }
        with (
            patch.object(settings, "backend_auth_enabled", True),
            patch.object(JobStore, "get", return_value=job),
        ):
            with self.assertRaises(HTTPException) as raised:
                publications_router._owned_export_job(
                    self._request("supabase-user-b"),
                    "job-a",
                )

        self.assertEqual(raised.exception.status_code, 404)

    def test_owner_can_read_own_export(self):
        job = {
            "id": "job-a",
            "job_type": "publication_export",
            "requested_by_user_id": "supabase-user-a",
        }
        with (
            patch.object(settings, "backend_auth_enabled", True),
            patch.object(JobStore, "get", return_value=job),
        ):
            result = publications_router._owned_export_job(
                self._request("supabase-user-a"),
                "job-a",
            )

        self.assertIs(result, job)

    def test_legacy_four_column_export_is_not_reused(self):
        legacy_job = {
            "id": "legacy-job",
            "status": "success",
            "total_rows": 2,
        }

        with (
            patch.object(
                publications_router.publication_export_store,
                "get_referenced_job_id",
                return_value="legacy-job",
            ),
            patch.object(JobStore, "get", return_value=legacy_job),
            patch.object(
                publications_router.publication_export_store,
                "release_reference",
            ) as release_reference,
        ):
            result = publications_router._existing_export_job(
                requested_by_user_id="supabase-user-a",
                seller_id="2682261950",
                creation_date="2026-07-24",
                total_rows=2,
            )

        self.assertIsNone(result)
        release_reference.assert_called_once_with(
            requested_by_user_id="supabase-user-a",
            seller_id="2682261950",
            creation_date="2026-07-24",
            job_id="legacy-job",
        )


class PublicationExportRecoveryTests(unittest.TestCase):
    def test_stale_processing_job_is_requeued(self):
        job = {
            "id": "job-stale",
            "status": "processing",
            "heartbeat_at": 100.0,
            "ml_user_id": "2682261950",
            "publication_date": "2026-06-17",
            "recovery_count": 0,
        }
        recovered_job = {
            **job,
            "status": "queued",
            "recovery_count": 1,
        }
        async_result = Mock(id="task-recovery")

        with (
            patch(
                "routers.publications_router.time.time",
                return_value=500.0,
            ),
            patch.object(
                settings,
                "ml_publication_export_recovery_stale_seconds",
                180,
            ),
            patch.object(
                publications_router.publication_export_store,
                "has_run_lock",
                return_value=False,
            ),
            patch.object(
                publications_router.publication_export_store,
                "try_acquire_recovery_guard",
                return_value=True,
            ),
            patch.object(
                publications_router.export_publications_task,
                "delay",
                return_value=async_result,
            ) as delay,
            patch.object(JobStore, "update", new=Mock()) as update,
            patch.object(JobStore, "get", return_value=recovered_job),
        ):
            result = publications_router._recover_stale_export_if_needed(job)

        delay.assert_called_once_with(
            "job-stale",
            "2682261950",
            "2026-06-17",
        )
        self.assertEqual(update.call_args.kwargs["recovery_count"], 1)
        self.assertEqual(result["status"], "queued")

    def test_active_run_lock_prevents_duplicate_recovery(self):
        job = {
            "id": "job-active",
            "status": "processing",
            "heartbeat_at": 100.0,
        }

        with (
            patch(
                "routers.publications_router.time.time",
                return_value=500.0,
            ),
            patch.object(
                settings,
                "ml_publication_export_recovery_stale_seconds",
                180,
            ),
            patch.object(
                publications_router.publication_export_store,
                "has_run_lock",
                return_value=True,
            ),
            patch.object(
                publications_router.export_publications_task,
                "delay",
            ) as delay,
        ):
            result = publications_router._recover_stale_export_if_needed(job)

        delay.assert_not_called()
        self.assertIs(result, job)


if __name__ == "__main__":
    unittest.main()
