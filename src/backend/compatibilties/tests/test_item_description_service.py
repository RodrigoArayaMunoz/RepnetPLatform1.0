import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, Mock, patch

import httpx
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from config import settings
from routers.process_queue_router import export_process_queue_result
from services.compatibility_service import JobMetrics
from services.item_description_service import (
    _process_item_description_row,
    build_item_description_excel,
    load_item_description_rows,
    process_item_description_job,
)
from services.ml_client import MercadoLibreClient
from services.process_queue_service import (
    _execute_process_record,
    _has_partial_process_errors,
    detect_process_type,
    run_process_queue,
)
from services.redis_rate_limiter import CombinedRateLimiter


def write_excel(path, rows, sheet="Hoja1", headers=("MLC ", "DESCRIPCION A ")):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    other = workbook.create_sheet("Hoja2")
    other.append(["MLC", "DESCRIPCION A"])
    other.append(["MLC999", "No procesar esta hoja"])
    workbook.save(path)
    workbook.close()


class DescriptionWorkbookTests(unittest.TestCase):
    def test_sheet_headers_unicode_line_breaks_and_formula_like_text_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "input.xlsx")
            text = ' "Descripción ñ\n\n12V\n" '
            write_excel(path, [(" mlc123 ", text), (None, None), ("MLC124", "=texto")])
            # Excel interprets a leading '=' as a formula unless the cell is text.
            workbook = load_workbook(path)
            workbook["Hoja1"]["B4"].data_type = "s"
            workbook.save(path)
            workbook.close()
            self.assertEqual(detect_process_type(str(path)), "item_descriptions")
            rows = load_item_description_rows(str(path))
        self.assertEqual([row["item_id"] for row in rows], ["MLC123", "MLC124"])
        self.assertEqual(rows[0]["plain_text"], text)
        self.assertEqual(rows[1]["plain_text"], "=texto")
        self.assertEqual(rows[1]["excel_row"], 4)

    def test_invalid_rows_and_conflicting_duplicates_are_rejected_before_writes(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "input.xlsx")
            write_excel(path, [("MLC1", "A"), ("MLC1", "B"), ("../items", "X"), ("MLC2", "  ")])
            rows = load_item_description_rows(str(path))
        self.assertTrue(all(row.get("validation_error") for row in rows))

    def test_missing_sheet_and_ambiguous_header_fail(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "input.xlsx")
            write_excel(path, [("MLC1", "A")], sheet="Datos")
            with self.assertRaisesRegex(ValueError, "Hoja1"):
                load_item_description_rows(str(path))
            write_excel(path, [("MLC1", "A", "B")], headers=("MLC", "DESCRIPCION A", "DESCRIPCIÓN A "))
            with self.assertRaisesRegex(ValueError, "exactamente una"):
                load_item_description_rows(str(path))

    def test_result_excel_is_reusable_and_does_not_execute_text_as_formulas(self):
        output = build_item_description_excel([
            {"item_id": "MLC1", "plain_text": "=1+1\nñ", "ok": True,
             "action": "updated", "reason": "Correcta", "excel_row": 2},
            {"item_id": "MLC2", "plain_text": "Texto", "ok": False,
             "reason": "Error de validación", "excel_row": 3},
        ])
        workbook = load_workbook(output)
        self.assertEqual(workbook.sheetnames, ["Hoja1"])
        self.assertEqual(workbook["Hoja1"]["B2"].data_type, "s")
        self.assertEqual(workbook["Hoja1"]["B2"].value, "=1+1\nñ")
        self.assertEqual(workbook["Hoja1"]["C3"].value, "Error")
        workbook.close()


class DescriptionRequestTests(unittest.IsolatedAsyncioTestCase):
    async def run_row(self, outcomes, attempts=3):
        calls = []
        def respond(request):
            calls.append(request)
            result = outcomes.pop(0)
            if isinstance(result, Exception):
                raise result
            code, body, headers = result
            return httpx.Response(code, json=body, headers=headers)
        client = MercadoLibreClient()
        client.get_valid_token = AsyncMock(return_value="test-token")
        client.client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        endpoint = Mock(window_seconds=60)
        endpoint.acquire = AsyncMock()
        endpoint.penalize = AsyncMock()
        global_write = Mock(window_seconds=60)
        global_write.acquire = AsyncMock()
        global_write.penalize = AsyncMock()
        combined = CombinedRateLimiter(endpoint, global_write)
        metrics = JobMetrics()
        row = {"item_id": "MLC123", "plain_text": "Nueva descripción\nñ", "excel_row": 2}
        try:
            with (
                patch("services.item_description_service.ml_client", client),
                patch("services.item_description_service.ITEM_DESCRIPTION_RATE_LIMITER", endpoint),
                patch("services.item_description_service.ITEM_DESCRIPTION_WRITE_RATE_LIMITER", combined),
                patch("services.ml_client.asyncio.sleep", new=AsyncMock()),
                patch("services.ml_client.random.uniform", return_value=0),
                patch.object(settings, "ml_retry_attempts", attempts),
            ):
                result = await _process_item_description_row(row, "99", metrics)
        finally:
            await client.client.aclose()
        return result, calls, metrics, endpoint, global_write

    async def test_missing_description_uses_post_and_exact_plain_text_body(self):
        result, calls, metrics, endpoint, global_write = await self.run_row([
            (404, {"message": "Description not found"}, {}),
            (201, {"plain_text": "Nueva descripción\nñ"}, {}),
        ])
        self.assertTrue(result["ok"])
        self.assertEqual(result["action"], "created")
        self.assertEqual([call.method for call in calls], ["GET", "POST"])
        self.assertEqual(calls[1].url.path, "/items/MLC123/description")
        self.assertEqual(json.loads(calls[1].content), {"plain_text": "Nueva descripción\nñ"})
        self.assertEqual(endpoint.acquire.await_count, 2)
        self.assertEqual(global_write.acquire.await_count, 1)
        self.assertEqual(metrics.ml_requests, 2)
        self.assertEqual(metrics.ml_http_errors, 0)

    async def test_existing_description_uses_put_with_api_version_2(self):
        result, calls, _, _, _ = await self.run_row([
            (200, {"plain_text": "Anterior"}, {}),
            (200, {"plain_text": "Nueva descripción\nñ"}, {}),
        ])
        self.assertTrue(result["ok"])
        self.assertEqual([call.method for call in calls], ["GET", "PUT"])
        self.assertEqual(calls[1].url.params["api_version"], "2")

    async def test_unchanged_description_only_reads(self):
        result, calls, _, _, global_write = await self.run_row([
            (200, {"plain_text": "Nueva descripción\r\nñ"}, {}),
        ])
        self.assertEqual(result["action"], "unchanged")
        self.assertEqual(len(calls), 1)
        global_write.acquire.assert_not_awaited()

    async def test_429_retries_spend_both_budgets_and_penalize_shared_state(self):
        result, calls, metrics, endpoint, global_write = await self.run_row([
            (200, {"plain_text": "Anterior"}, {}),
            (429, {"message": "Too many requests"}, {"Retry-After": "90"}),
            (200, {"plain_text": "Nueva descripción\nñ"}, {}),
        ])
        self.assertTrue(result["ok"])
        self.assertEqual(metrics.ml_requests, 3)
        self.assertEqual(metrics.ml_retries, 1)
        self.assertEqual(metrics.ml_rate_limited, 1)
        self.assertEqual(endpoint.acquire.await_count, 3)
        self.assertEqual(global_write.acquire.await_count, 2)
        endpoint.penalize.assert_awaited_with(90)
        global_write.penalize.assert_awaited_with(90)

    async def test_metadata_write_response_is_verified_with_get(self):
        result, calls, _, _, _ = await self.run_row([
            (200, {"plain_text": "Anterior"}, {}),
            (200, {"last_updated": "2026-10-02"}, {}),
            (200, {"plain_text": "Nueva descripción\nñ"}, {}),
        ])
        self.assertTrue(result["ok"])
        self.assertEqual([call.method for call in calls], ["GET", "PUT", "GET"])

    async def test_read_429_blocks_the_endpoint_without_spending_write_budget(self):
        result, calls, metrics, endpoint, global_write = await self.run_row([
            (429, {"message": "Too many requests"}, {}),
            (200, {"plain_text": "Nueva descripción\nñ"}, {}),
        ])
        self.assertTrue(result["ok"])
        self.assertEqual([call.method for call in calls], ["GET", "GET"])
        self.assertEqual(metrics.ml_rate_limited, 1)
        global_write.acquire.assert_not_awaited()
        global_write.penalize.assert_not_awaited()
        endpoint.penalize.assert_awaited_with(60)

    async def test_unconfirmed_text_is_not_reported_as_success(self):
        result, _, _, _, _ = await self.run_row([
            (200, {"plain_text": "Anterior"}, {}),
            (200, {"plain_text": "Anterior"}, {}),
        ])
        self.assertFalse(result["ok"])
        self.assertEqual(result["error_code"], "DESCRIPTION_NOT_CONFIRMED")

    async def test_post_applied_before_timeout_is_recovered_after_duplicate_response(self):
        result, calls, metrics, _, _ = await self.run_row([
            (404, {}, {}),
            httpx.ReadTimeout("lost write response"),
            (400, {"message": "Description already exists"}, {}),
            (200, {"plain_text": "Nueva descripción\nñ"}, {}),
        ])
        self.assertTrue(result["ok"])
        self.assertEqual([call.method for call in calls], ["GET", "POST", "POST", "GET"])
        self.assertEqual(metrics.ml_retries, 1)
        self.assertEqual(metrics.ml_http_errors, 1)

    async def test_permission_error_does_not_attempt_a_write(self):
        result, calls, _, _, global_write = await self.run_row([(403, {"message": "Forbidden"}, {})])
        self.assertFalse(result["ok"])
        self.assertEqual(len(calls), 1)
        global_write.acquire.assert_not_awaited()

    async def test_invalid_row_never_calls_the_api(self):
        with patch("services.item_description_service.ml_client") as client:
            result = await _process_item_description_row({"validation_error": "Vacía"}, "99", JobMetrics())
        self.assertFalse(result["ok"])
        client.get_item_description.assert_not_called()


class DescriptionQueueTests(unittest.IsolatedAsyncioTestCase):
    async def test_non_text_description_is_reported_without_crashing_the_result_file(self):
        from datetime import datetime
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "input.xlsx")
            write_excel(path, [("MLC1", datetime(2026, 10, 2))])
            with (
                patch("services.item_description_service.ml_client") as client,
                patch("services.item_description_service.JobStore.update"),
                patch.object(settings, "upload_dir", directory),
            ):
                result = await process_item_description_job("invalid", "99", str(path))
            self.assertTrue(Path(directory, "invalid_item_descriptions_result.json").exists())
        self.assertEqual(result["summary"]["error_count"], 1)
        client.get_item_description.assert_not_called()

    async def test_real_excel_parser_job_concurrency_duplicates_and_order(self):
        active = maximum = 0
        calls = []
        async def process(row, user_id, metrics):
            nonlocal active, maximum
            active += 1
            maximum = max(maximum, active)
            calls.append(row["item_id"])
            await asyncio.sleep(0)
            active -= 1
            return {**row, "ok": True, "action": "updated", "reason": "Actualizada"}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "input.xlsx")
            write_excel(path, [("MLC1", "A"), ("MLC2", "B"), ("MLC1", "A"), ("MLC3", "C")])
            with (
                patch("services.item_description_service._process_item_description_row", side_effect=process),
                patch("services.item_description_service.JobStore.update") as updates,
                patch.object(settings, "upload_dir", directory),
                patch.object(settings, "item_description_chunk_size", 2),
                patch.object(settings, "item_description_max_concurrency", 2),
                patch.object(settings, "job_progress_update_every", 2),
            ):
                result = await process_item_description_job("test", "99", str(path))
            saved = json.loads(Path(directory, "test_item_descriptions_result.json").read_text())
        self.assertEqual(calls, ["MLC1", "MLC2", "MLC3"])
        self.assertEqual(maximum, 2)
        self.assertEqual([row["item_id"] for row in saved], ["MLC1", "MLC2", "MLC1", "MLC3"])
        self.assertEqual(result["summary"]["duplicate_rows"], 1)
        self.assertEqual(result["summary"]["updated_items"], 3)
        self.assertEqual(result["summary"]["success_count"], 4)
        self.assertFalse(any("Esperando" in str(call.kwargs) for call in updates.call_args_list))

    async def test_dispatch_and_partial_errors_use_the_new_process_type(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "input.xlsx")
            write_excel(path, [("MLC1", "A")])
            with (
                patch("services.process_queue_service.supabase_process_store.download_process_file", new=AsyncMock(return_value=str(path))),
                patch("services.process_queue_service.process_queue_store.update"),
                patch("services.process_queue_service._run_item_descriptions_job", new=AsyncMock(return_value=("job", {"error_count": 1}))) as run,
            ):
                process_type, job_id, summary = await _execute_process_record(process_row={"archivo": "input.xlsx"}, user_id="99")
        self.assertEqual((process_type, job_id), ("item_descriptions", "job"))
        self.assertTrue(_has_partial_process_errors(process_type, summary))
        run.assert_awaited_once()

    async def test_result_export_route_returns_the_description_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "result.json")
            path.write_text(json.dumps([{"item_id": "MLC1", "plain_text": "Texto", "ok": True, "action": "created"}]))
            with patch("routers.process_queue_router.process_queue_result_store.get", return_value={"process_type": "item_descriptions", "result_path": str(path)}):
                response = await export_process_queue_result("record")
                content = b"".join([chunk if isinstance(chunk, bytes) else chunk.encode() async for chunk in response.body_iterator])
        self.assertIn("resultado_descripciones_mlc", response.headers["content-disposition"])
        from io import BytesIO
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook["Hoja1"]["A2"].value, "MLC1")
        workbook.close()

    async def test_queue_waits_five_minutes_after_a_description_file_and_exports_its_result(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory, "result.json")
            path.write_text("[]")
            first = {"id": 1, "archivo": "descripciones.xlsx"}
            second = {"id": 2, "archivo": "fotos.xlsx"}
            store = Mock()
            store.can_access = True
            store.list_pending_processes = AsyncMock(side_effect=[[first, second], [second], [second], []])
            store.update_process_status = AsyncMock()
            with (
                patch("services.process_queue_service.supabase_process_store", store),
                patch("services.process_queue_service.ml_client.startup", new=AsyncMock()),
                patch("services.process_queue_service.ml_client.shutdown", new=AsyncMock()),
                patch("services.process_queue_service.supabase_meli_connection_store.restore_token_store", new=AsyncMock()),
                patch("services.process_queue_service._execute_process_record", new=AsyncMock(side_effect=[("item_descriptions", "a", {"error_count": 0, "exported_rows_total": 1}), ("item_pictures", "b", {"error_count": 0})])),
                patch("services.process_queue_service.JobStore.get", return_value={"result_path": str(path)}),
                patch("services.process_queue_service.process_queue_store"),
                patch("services.process_queue_service.process_queue_error_store"),
                patch("services.process_queue_service.process_queue_result_store") as results,
                patch("services.process_queue_service.asyncio.sleep", new=AsyncMock()) as sleep,
            ):
                await run_process_queue(user_id="99")
        sleep.assert_awaited_once_with(300)
        self.assertEqual(results.save.call_args.args[1]["export_kind"], "item_descriptions")


if __name__ == "__main__":
    unittest.main()
