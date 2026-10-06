import base64
import copy
import json
import unittest
from datetime import UTC, datetime, timedelta
from io import BytesIO
from unittest.mock import patch

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook

from config import settings
from routers.refax_router import router
from services.refax_connection_service import RefaxConnectionService
from services.refax_products_service import RefaxProductsService, XLSX_CONTENT_TYPE
from services.supabase_refax_connection_store import SupabaseRefaxConnectionStore


class RefaxEnvironmentConnectionTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.now = datetime(2026, 10, 5, 15, 0, tzinfo=UTC)
        configuration = {
            "supabase_url": "https://supabase.example.test",
            "supabase_service_role_key": "supabase-private-key",
            "refax_api_base_url": "https://production.refax.example.test",
            "refax_provider_code": "production-code",
            "refax_api_key": "production-private-key",
            "refax_country_code": 1,
            "refax_test_api_base_url": "http://apitest.refax.com:8098",
            "refax_test_provider_code": "test-code",
            "refax_test_api_key": "test-private-key#@",
            "refax_test_country_code": 1,
        }
        for name, value in configuration.items():
            self.enterContext(patch.object(settings, name, value))
        self.rows = {
            1: {
                "id": 1, "environment": "production", "is_active": True,
                "access_token": "original-production-token", "last_error": None,
                "refresh_at": (self.now + timedelta(hours=1)).isoformat(),
                "expires_at": (self.now + timedelta(hours=2)).isoformat(),
            },
            2: {
                "id": 2, "environment": "test", "is_active": False,
                "access_token": None, "last_error": None,
            },
        }
        self.original_production = copy.deepcopy(self.rows[1])
        self.auth_requests = []
        self.product_requests = []
        self.product_statuses = []
        self.products = [{
            "numero_refax": "000123", "precio": 1500.0, "stock": "Disponible",
            "nombre_producto": "ESPEJO TEST", "glosa_1": "A", "glosa_2": "B", "glosa_3": "C",
            "factor": 1, "marca_producto": "MARCA", "origen": "ORIGEN", "rubro_comercial": "RUBRO",
            "imagenes_url": ["https://example.test/a.jpg", "https://example.test/b.jpg"],
            "codigo_oem": ["001", "002"], "marca_oem": ["MARCA A", "MARCA B"],
            "codigo_fabrica": ["F01", "F02"], "fabrica": ["FABRICA A", "FABRICA B"],
            "aplicaciones": ["APLICACION A", "APLICACION B"],
            "precio_oferta_web": 19.95, "cantidad_oferta_web": 0,
        }]
        self.test_products_payload = copy.deepcopy(self.products)
        for product in self.test_products_payload:
            product["aplicacion"] = product.pop("aplicaciones")
            product["imagen_url"] = product.pop("imagenes_url")
        self.reject_test_auth = False
        original_client = httpx.AsyncClient
        self.original_client = original_client
        transport = httpx.MockTransport(self.handle_request)
        self.enterContext(patch(
            "httpx.AsyncClient",
            side_effect=lambda **kwargs: original_client(transport=transport, **kwargs),
        ))
        self.production = RefaxConnectionService()
        self.test = RefaxConnectionService(environment="test")
        self.production._now = lambda: self.now
        self.test._now = lambda: self.now
        self.enterContext(patch("routers.refax_router.refax_connection_service", self.production))
        self.enterContext(patch("routers.refax_router.refax_test_connection_service", self.test))
        self.production_products = RefaxProductsService(self.production)
        self.test_products = RefaxProductsService(self.test, environment="test")
        self.enterContext(patch("routers.refax_router.refax_products_service", self.production_products))
        self.enterContext(patch("routers.refax_router.refax_test_products_service", self.test_products))
        self.app = FastAPI()
        self.app.include_router(router)

    def handle_request(self, request):
        if request.url.host in ("apitest.refax.com", "production.refax.example.test"):
            environment = "test" if request.url.host == "apitest.refax.com" else "production"
            if request.url.path == "/api/Productos/Listado":
                self.assertEqual(request.method, "GET")
                self.product_requests.append(request)
                status = self.product_statuses.pop(0) if self.product_statuses else 200
                products = self.test_products_payload if environment == "test" else self.products
                return httpx.Response(status, json=products if status == 200 else {"error": "private-upstream-details"})
            self.auth_requests.append((str(request.url), json.loads(request.content)))
            if environment == "test" and self.reject_test_auth:
                return httpx.Response(401)
            return httpx.Response(200, json={"status": "OK", "token": f"new-{environment}-token"})

        self.assertEqual(request.url.host, "supabase.example.test")
        if request.method == "POST":
            payload = json.loads(request.content)
            connection_id = payload["id"]
            self.assertEqual(payload["environment"], self.rows[connection_id]["environment"])
            self.rows[connection_id].update(payload)
            return httpx.Response(200, json=[self.rows[connection_id]])

        connection_id = int(request.url.params["id"].removeprefix("eq."))
        self.assertEqual(
            request.url.params["environment"], f"eq.{self.rows[connection_id]['environment']}"
        )
        if request.method == "GET":
            return httpx.Response(200, json=[self.rows[connection_id]])
        if request.method == "PATCH":
            self.rows[connection_id].update(json.loads(request.content))
            return httpx.Response(204)
        raise AssertionError("Unexpected request method")

    async def request(self, method, path):
        async with self.original_client(
            transport=httpx.ASGITransport(app=self.app), base_url="http://backend.example.test"
        ) as client:
            return await client.request(method, path)

    async def test_test_connect_persists_token_and_restores_status_without_changing_production(self):
        response = await self.request("POST", "/refax/test/connect")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["connected"])
        self.assertNotIn("access_token", response.json())
        self.assertEqual(self.auth_requests, [(
            "http://apitest.refax.com:8098/api/Autenticacion/GetToken",
            {"codigo": "test-code", "clave": "test-private-key#@", "pais": 1},
        )])
        self.assertEqual(self.rows[2]["access_token"], "new-test-token")
        self.assertEqual(self.rows[2]["obtained_at"], self.now.isoformat())
        self.assertEqual(self.rows[2]["refresh_at"], (self.now + timedelta(hours=7, minutes=45)).isoformat())
        self.assertEqual(self.rows[2]["expires_at"], (self.now + timedelta(hours=8)).isoformat())
        restored = await self.request("GET", "/refax/test/status")
        self.assertTrue(restored.json()["connected"])
        self.assertEqual(self.rows[1], self.original_production)

    async def test_test_authentication_failure_updates_only_test(self):
        self.reject_test_auth = True
        response = await self.request("POST", "/refax/test/connect")
        self.assertEqual(response.status_code, 502)
        self.assertFalse(self.rows[2]["is_active"])
        self.assertIn("HTTP 401", self.rows[2]["last_error"])
        self.assertEqual(self.rows[1], self.original_production)

    async def test_missing_test_credentials_never_fall_back_to_production(self):
        with patch.object(settings, "refax_test_api_key", None):
            response = await self.request("POST", "/refax/test/connect")
        self.assertEqual(response.status_code, 502)
        self.assertIn("REFAX_TEST_API_KEY", response.json()["detail"])
        self.assertEqual(self.auth_requests, [])
        self.assertEqual(self.rows[1], self.original_production)

    async def test_due_test_token_is_renewed_with_test_credentials(self):
        self.rows[2].update({
            "is_active": True,
            "access_token": "old-test-token",
            "refresh_at": (self.now - timedelta(minutes=1)).isoformat(),
            "expires_at": (self.now + timedelta(minutes=10)).isoformat(),
        })
        response = await self.request("GET", "/refax/test/status")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["connected"])
        self.assertEqual(self.rows[2]["access_token"], "new-test-token")
        self.assertEqual(self.auth_requests[0][1]["codigo"], "test-code")
        self.assertEqual(self.rows[1], self.original_production)

    async def test_production_connect_keeps_original_endpoint_and_credentials(self):
        original_test = copy.deepcopy(self.rows[2])
        response = await self.request("POST", "/refax/connect")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["connected"])
        self.assertEqual(self.auth_requests, [(
            "https://production.refax.example.test/api/Autenticacion/GetToken",
            {"codigo": "production-code", "clave": "production-private-key", "pais": 1},
        )])
        self.assertEqual(self.rows[2], original_test)

    def test_unknown_environment_is_rejected(self):
        with self.assertRaises(ValueError):
            SupabaseRefaxConnectionStore("invalid")
        with self.assertRaises(ValueError):
            RefaxConnectionService(environment="invalid")
        with self.assertRaises(ValueError):
            RefaxProductsService(environment="invalid")

    def connect_test_row(self):
        self.rows[2].update({
            "is_active": True, "access_token": "original-test-token",
            "refresh_at": (self.now + timedelta(hours=1)).isoformat(),
            "expires_at": (self.now + timedelta(hours=2)).isoformat(),
        })

    def assert_test_request(self, request, token):
        self.assertEqual(
            str(request.url), "http://apitest.refax.com:8098/api/Productos/Listado?codigo=test-code"
        )
        self.assertEqual(request.headers["authorization"], f"Bearer {token}")

    def assert_excel_matches_production(self, content):
        expected = RefaxProductsService._build_excel(json.dumps(self.products).encode())
        actual_workbook = load_workbook(BytesIO(content))
        expected_workbook = load_workbook(BytesIO(expected))
        try:
            actual = actual_workbook["Productos"]
            production = expected_workbook["Productos"]
            self.assertEqual(actual.max_column, 19)
            self.assertEqual(list(actual.values), list(production.values))
            self.assertEqual(actual["A2"].value, "000123")
            self.assertEqual(actual["L2"].value.splitlines(), self.products[0]["imagenes_url"])
            self.assertEqual(actual["Q2"].value.splitlines(), self.products[0]["aplicaciones"])
            self.assertEqual(actual.freeze_panes, production.freeze_panes)
            self.assertEqual(actual.auto_filter.ref, production.auto_filter.ref)
            for actual_row, production_row in zip(actual.iter_rows(), production.iter_rows()):
                for cell, expected_cell in zip(actual_row, production_row):
                    self.assertEqual(cell.data_type, expected_cell.data_type)
                    self.assertEqual(cell.number_format, expected_cell.number_format)
                    self.assertEqual(copy.copy(cell.alignment), copy.copy(expected_cell.alignment))
        finally:
            actual_workbook.close()
            expected_workbook.close()

    async def test_test_download_uses_persisted_test_token_and_exports_same_excel_as_production(self):
        self.connect_test_row()
        response = await self.request("GET", "/refax/test/products/download")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], XLSX_CONTENT_TYPE)
        self.assertIn('filename="productos_refax_test_', response.headers["content-disposition"])
        self.assertEqual(len(self.product_requests), 1)
        self.assert_test_request(self.product_requests[0], "original-test-token")
        self.assert_excel_matches_production(response.content)
        self.assertEqual(self.auth_requests, [])
        self.assertEqual(self.rows[1], self.original_production)

    async def test_test_stream_exports_same_excel_and_progress_without_exposing_tokens(self):
        self.connect_test_row()
        response = await self.request("GET", "/refax/test/products/download?progress=true")
        self.assertEqual(response.headers["content-type"], "application/x-ndjson")
        self.assertEqual(response.headers["x-accel-buffering"], "no")
        events = [json.loads(line) for line in response.text.splitlines()]
        self.assertEqual(events[-1]["type"], "complete")
        metadata = next(event for event in events if event["type"] == "file")
        self.assertTrue(metadata["filename"].startswith("productos_refax_test_"))
        content = b"".join(base64.b64decode(event["data"]) for event in events if event["type"] == "chunk")
        self.assertEqual(metadata["size"], len(content))
        self.assert_excel_matches_production(content)
        self.assert_test_request(self.product_requests[0], "original-test-token")
        self.assertNotIn("original-test-token", response.text)
        self.assertNotIn("original-production-token", response.text)
        self.assertEqual(self.rows[1], self.original_production)

    async def test_disconnected_test_does_not_query_products_or_use_production(self):
        response = await self.request("GET", "/refax/test/products/download")
        self.assertEqual(response.status_code, 502)
        self.assertIn("no esta conectado", response.json()["detail"])
        self.assertEqual(self.product_requests, [])
        streamed = await self.request("GET", "/refax/test/products/download?progress=true")
        events = [json.loads(line) for line in streamed.text.splitlines()]
        self.assertEqual(events[-1]["type"], "error")
        self.assertFalse(any(event["type"] == "file" for event in events))
        self.assertEqual(self.rows[1], self.original_production)

    async def test_rejected_test_token_renews_only_test_before_retrying_products(self):
        self.connect_test_row()
        self.product_statuses = [401, 200]
        response = await self.request("GET", "/refax/test/products/download")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(self.product_requests), 2)
        self.assert_test_request(self.product_requests[0], "original-test-token")
        self.assert_test_request(self.product_requests[1], "new-test-token")
        self.assertEqual(len(self.auth_requests), 1)
        self.assertEqual(self.auth_requests[0][1]["codigo"], "test-code")
        self.assertEqual(self.rows[1], self.original_production)

    async def test_missing_test_provider_code_does_not_fall_back_to_production(self):
        self.connect_test_row()
        with patch.object(settings, "refax_test_provider_code", None):
            response = await self.request("GET", "/refax/test/products/download")
        self.assertEqual(response.status_code, 502)
        self.assertIn("REFAX_TEST_PROVIDER_CODE", response.json()["detail"])
        self.assertEqual(self.product_requests, [])
        self.assertEqual(self.rows[1], self.original_production)

    async def test_production_download_keeps_its_endpoint_code_and_token(self):
        original_test = copy.deepcopy(self.rows[2])
        response = await self.request("GET", "/refax/products/download")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            str(self.product_requests[0].url),
            "https://production.refax.example.test/api/Productos/Listado?codigo=production-code",
        )
        self.assertEqual(self.product_requests[0].headers["authorization"], "Bearer original-production-token")
        self.assertNotIn("productos_refax_test_", response.headers["content-disposition"])
        self.assertEqual(self.rows[2], original_test)


if __name__ == "__main__":
    unittest.main()
