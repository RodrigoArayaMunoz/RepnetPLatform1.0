import json
import unittest
from io import BytesIO
from unittest.mock import AsyncMock, patch
from xml.etree import ElementTree
from zipfile import ZipFile

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook

from config import settings
from routers.refax_router import router
from services.refax_products_service import (
    RefaxProductsError,
    RefaxProductsService,
    XLSX_CONTENT_TYPE,
)

EXPECTED_HEADERS = (
    "SKU", "PRECIO", "STOCK", "NOMBRE_PRODUCTO", "GLOSA_1", "GLOSA_2",
    "GLOSA_3", "FACTOR", "MARCA_PRODUCTO", "ORIGEN", "RUBRO_COMERCIAL",
    "IMAGENES_URL", "CODIGO_OEM", "MARCA_OEM", "CODIGO_FABRICA", "FABRICA",
    "APLICACIONES", "PRECIO_OFERTA_WEB", "CANTIDAD_OFERTA_WEB",
)


class RefaxProductsServiceTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.enterContext(patch.object(settings, "refax_provider_code", "test-provider"))
        self.enterContext(patch.object(settings, "refax_api_base_url", "https://refax.example.test"))

    def build_service(self):
        connection_service = AsyncMock()
        connection_service.get_valid_token.return_value = "private-token"
        service = RefaxProductsService(connection_service)
        return service, connection_service

    async def test_download_preserves_text_format_and_adds_product_columns(self):
        service, connection_service = self.build_service()
        service._request_products = AsyncMock(
            return_value=httpx.Response(
                200,
                json=[
                    {"numero_refax": "00123", "precio": 1500, "stock": "5", "nombre_producto": "ESPEJO EXTERIOR"},
                    {"numero_refax": "ABC", "precio": 19.95, "stock": "Disponible"},
                    {"numero_refax": "=1+1", "precio": 0, "stock": "0"},
                    {"numero_refax": "00000", "precio": "1250.00", "stock": 6},
                ],
                headers={"content-type": "application/json; charset=utf-8"},
            )
        )

        result = await service.download()

        self.assertEqual(result.content_type, XLSX_CONTENT_TYPE)
        workbook = load_workbook(BytesIO(result.content))
        try:
            sheet = workbook["Productos"]
            self.assertEqual(list(sheet.values), [
                EXPECTED_HEADERS,
                ("00123", "1500", "5", "ESPEJO EXTERIOR") + (None,) * 15,
                ("ABC", "19.95", "Disponible") + (None,) * 16,
                ("=1+1", "0", "0") + (None,) * 16,
                ("00000", "1250", "6") + (None,) * 16,
            ])
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if cell.value is not None:
                        self.assertEqual(cell.data_type, "s")
                    self.assertEqual(cell.number_format, "@")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:S5")
        finally:
            workbook.close()
        # Excel consulta esta regla, no basta con aplicar el formato Texto.
        with ZipFile(BytesIO(result.content)) as archive:
            self.assertIsNone(archive.testzip())
            sheet_xml = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            ignored = sheet_xml.find("s:ignoredErrors/s:ignoredError", namespace)
            self.assertIsNotNone(ignored)
            self.assertEqual(ignored.attrib, {"sqref": "A2:S5", "numberStoredAsText": "1"})
        connection_service.get_valid_token.assert_awaited_once()
        service._request_products.assert_awaited_once_with("private-token")

    async def test_download_exports_all_fields_from_refax_example(self):
        images = [
            f"https://imagenes.refaxchile.cl:9092/FOTOGRAFIAS/0000001/0000001{suffix}.jpg"
            for suffix in "ABCD"
        ]
        application = "CHEVROLET S10 PICK UP 2200 134CID L4 SOHC 8 VALV [1994 - 1998]"
        product = {
            "numero_refax": "0000001", "stock": "DISPONIBLE", "precio": 13990,
            "nombre_producto": "ESPEJO EXTERIOR", "glosa_1": "IZQUIERDO ELECTRICO NEGRO",
            "glosa_2": "", "glosa_3": "", "factor": 1, "marca_producto": "TYC",
            "origen": "TAIWAN", "rubro_comercial": "ESPEJOS", "imagenes_url": images,
            "codigo_oem": "17801665", "marca_oem": "", "codigo_fabrica": "388-GMD016",
            "fabrica": "TYC", "aplicaciones": application,
            "precio_oferta_web": 13990.00, "cantidad_oferta_web": 1,
        }
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=[product]))

        result = await service.download()

        workbook = load_workbook(BytesIO(result.content))
        try:
            sheet = workbook["Productos"]
            self.assertEqual(list(sheet.values), [
                EXPECTED_HEADERS,
                ("0000001", "13990", "DISPONIBLE", "ESPEJO EXTERIOR",
                 "IZQUIERDO ELECTRICO NEGRO", None, None, "1", "TYC", "TAIWAN",
                 "ESPEJOS", "\n".join(images), "17801665", None, "388-GMD016",
                 "TYC", application, "13990", "1"),
            ])
            self.assertTrue(sheet["L2"].alignment.wrap_text)
            self.assertTrue(sheet["Q2"].alignment.wrap_text)
        finally:
            workbook.close()

    async def test_applications_preserve_single_text_multiple_texts_and_json_nodes(self):
        applications = [
            "CHEVROLET S10 [1994 - 1998]",
            "CHEVROLET S10 [1994 - 1998]\nCHEVROLET BLAZER [1995 - 2000]",
            ["CHEVROLET S10 [1994 - 1998]", "CHEVROLET BLAZER [1995 - 2000]"],
            [
                {"marca": "CHEVROLET", "modelo": "S10", "años": [1994, 1998]},
                {"marca": "CHEVROLET", "modelo": "BLAZER", "años": [1995, 2000]},
            ],
            {"marca": "CHEVROLET", "modelos": ["S10", "BLAZER"]},
            [], None,
        ]
        products = [
            {"numero_refax": str(index).zfill(7), "precio": 0, "stock": "DISPONIBLE",
             "aplicaciones": value, "factor": 0, "precio_oferta_web": 0,
             "cantidad_oferta_web": 0, "imagenes_url": []}
            for index, value in enumerate(applications)
        ]
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=products))

        result = await service.download()

        workbook = load_workbook(BytesIO(result.content))
        try:
            rows = list(workbook.active.values)[1:]
            self.assertEqual(rows[0][16], applications[0])
            self.assertEqual(rows[1][16], applications[1])
            self.assertEqual(rows[2][16].splitlines(), applications[2])
            self.assertEqual([json.loads(line) for line in rows[3][16].splitlines()], applications[3])
            self.assertEqual(json.loads(rows[4][16]), applications[4])
            self.assertIsNone(rows[5][16])
            self.assertIsNone(rows[6][16])
            for row in rows:
                self.assertEqual((row[7], row[17], row[18]), ("0", "0", "0"))
                self.assertIsNone(row[11])
        finally:
            workbook.close()

    async def test_singular_test_fields_preserve_single_and_multiple_applications_and_images(self):
        products = [
            {
                "numero_refax": "0000001", "precio": 1500, "stock": "Disponible",
                "aplicacion": "CHEVROLET S10 [1994 - 1998]",
                "imagen_url": ["https://example.test/a.jpg", "https://example.test/b.jpg"],
            },
            {
                "numero_refax": "0000002", "precio": 0, "stock": 0,
                "aplicacion": ["CHEVROLET S10 [1994 - 1998]", "CHEVROLET COMBO [2003 - 2005]"],
                "imagen_url": "https://example.test/c.jpg",
            },
        ]
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=products))
        result = await service.download()
        workbook = load_workbook(BytesIO(result.content))
        try:
            sheet = workbook["Productos"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), EXPECTED_HEADERS)
            self.assertEqual(sheet["Q2"].value, products[0]["aplicacion"])
            self.assertEqual(sheet["Q3"].value.splitlines(), products[1]["aplicacion"])
            self.assertEqual(sheet["L2"].value.splitlines(), products[0]["imagen_url"])
            self.assertEqual(sheet["L3"].value, products[1]["imagen_url"])
            self.assertEqual(sheet["B3"].value, "0")
            self.assertEqual(sheet["C3"].value, "0")
            for coordinate in ("L2", "L3", "Q2", "Q3"):
                self.assertEqual(sheet[coordinate].data_type, "s")
                self.assertEqual(sheet[coordinate].number_format, "@")
                self.assertTrue(sheet[coordinate].alignment.wrap_text)
        finally:
            workbook.close()

    async def test_aliases_preserve_plural_values_and_use_singular_when_plural_is_empty(self):
        products = [
            {
                "numero_refax": str(index).zfill(7), "precio": 0, "stock": "Disponible",
                "aplicaciones": value, "aplicacion": "APPLICATION FROM TEST",
                "imagenes_url": value, "imagen_url": "https://example.test/test.jpg",
            }
            for index, value in enumerate((None, "", [], "CANONICAL VALUE"))
        ]
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=products))
        result = await service.download()
        workbook = load_workbook(BytesIO(result.content))
        try:
            rows = list(workbook["Productos"].values)[1:]
            for row in rows[:3]:
                self.assertEqual(row[16], "APPLICATION FROM TEST")
                self.assertEqual(row[11], "https://example.test/test.jpg")
            self.assertEqual(rows[3][16], "CANONICAL VALUE")
            self.assertEqual(rows[3][11], "CANONICAL VALUE")
        finally:
            workbook.close()

    async def test_long_singular_application_fails_instead_of_being_silently_truncated(self):
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=[{
            "numero_refax": "0000001", "precio": 0, "stock": "Disponible",
            "aplicacion": "x" * 32768,
        }]))
        with self.assertRaisesRegex(RefaxProductsError, "APLICACIONES.*32.767"):
            await service.download()

    async def test_additional_fields_are_text_and_never_excel_formulas(self):
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=[{
            "numero_refax": "0000001", "precio": 1, "stock": "DISPONIBLE",
            "nombre_producto": "=1+1", "codigo_oem": "0000123",
            "aplicaciones": ["=1+1", "+2"],
        }]))

        result = await service.download()

        workbook = load_workbook(BytesIO(result.content))
        try:
            self.assertEqual(workbook.active["D2"].value, "=1+1")
            self.assertEqual(workbook.active["M2"].value, "0000123")
            self.assertEqual(workbook.active["Q2"].value, "=1+1\n+2")
            for column in ("D", "M", "Q"):
                self.assertEqual(workbook.active[f"{column}2"].data_type, "s")
        finally:
            workbook.close()

    async def test_download_exports_oem_factory_codes_and_factories_as_lists(self):
        product = {
            "numero_refax": "0000006", "precio": 13990, "stock": "DISPONIBLE",
            "codigo_oem": ["17540-85E00", "4708770", "51821653"],
            "codigo_fabrica": ["0-N1526", "27365", "534-0053-10"],
            "fabrica": ["OPTIMAL", "FEBI BILSTEIN", "INA"],
            "aplicaciones": ["CHEVROLET COMBO VAN 1300 Z13DT DOHC 16 VALV [2005 - 2011]"],
        }
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=[product]))

        result = await service.download()

        workbook = load_workbook(BytesIO(result.content))
        try:
            sheet = workbook["Productos"]
            self.assertEqual(sheet["M2"].value.splitlines(), product["codigo_oem"])
            self.assertEqual(sheet["O2"].value.splitlines(), product["codigo_fabrica"])
            self.assertEqual(sheet["P2"].value.splitlines(), product["fabrica"])
            self.assertEqual(sheet["A2"].value, "0000006")
        finally:
            workbook.close()

    async def test_every_column_accepts_lists_including_prices_stock_and_sku(self):
        product = {
            "numero_refax": ["0000006"], "precio": [13990.00, 0, "19.95"],
            "stock": ["DISPONIBLE", 0], "nombre_producto": ["ESPEJO", "=1+1"],
            "glosa_1": ["IZQUIERDO", "NEGRO"], "glosa_2": ["A", "B"],
            "glosa_3": ["C", "D"], "factor": [1, 0],
            "marca_producto": ["TYC", "INA"], "origen": ["TAIWAN", "CHINA"],
            "rubro_comercial": ["ESPEJOS", "ACCESORIOS"],
            "imagenes_url": ["https://example.test/A.jpg", "https://example.test/B.jpg"],
            "codigo_oem": ["00001", "00002"], "marca_oem": ["GM", "OPEL"],
            "codigo_fabrica": ["0-N1526", "27365"], "fabrica": ["OPTIMAL", "INA"],
            "aplicaciones": ["CHEVROLET COMBO", "OPEL CORSA"],
            "precio_oferta_web": ["1250.00", 0], "cantidad_oferta_web": [1, 2],
        }
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=[product]))

        result = await service.download()

        workbook = load_workbook(BytesIO(result.content))
        try:
            sheet = workbook["Productos"]
            self.assertEqual(list(sheet.values), [
                EXPECTED_HEADERS,
                ("0000006", "13990\n0\n19.95", "DISPONIBLE\n0", "ESPEJO\n=1+1",
                 "IZQUIERDO\nNEGRO", "A\nB", "C\nD", "1\n0", "TYC\nINA",
                 "TAIWAN\nCHINA", "ESPEJOS\nACCESORIOS",
                 "https://example.test/A.jpg\nhttps://example.test/B.jpg",
                 "00001\n00002", "GM\nOPEL", "0-N1526\n27365", "OPTIMAL\nINA",
                 "CHEVROLET COMBO\nOPEL CORSA", "1250\n0", "1\n2"),
            ])
            for cell in sheet[2]:
                self.assertEqual(cell.data_type, "s")
                self.assertEqual(cell.number_format, "@")
                self.assertTrue(cell.alignment.wrap_text)
        finally:
            workbook.close()

    async def test_optional_fields_preserve_nested_nodes_empty_values_and_numeric_zero(self):
        nested = {"fabricantes": ["INA", "OPTIMAL"], "codigo": "00001", "activo": False}
        products = [
            {"numero_refax": "0000006", "precio": 0, "stock": "DISPONIBLE",
             "codigo_oem": ["00001", nested, ["00002", "00003"], None, ""],
             "fabrica": nested, "glosa_1": False,
             "factor": [None, "", 0, 1.5], "precio_oferta_web": [],
             "cantidad_oferta_web": None},
            {"numero_refax": "0000007", "precio": 0, "stock": "DISPONIBLE",
             "codigo_oem": [], "codigo_fabrica": None, "fabrica": []},
        ]
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=products))

        result = await service.download()

        workbook = load_workbook(BytesIO(result.content))
        try:
            sheet = workbook["Productos"]
            entries = sheet["M2"].value.split("\n")
            self.assertEqual(entries[0], "00001")
            self.assertEqual(json.loads(entries[1]), nested)
            self.assertEqual(json.loads(entries[2]), ["00002", "00003"])
            self.assertEqual(entries[3:], ["", ""])
            self.assertEqual(json.loads(sheet["P2"].value), nested)
            self.assertEqual(sheet["E2"].value, "false")
            self.assertEqual(sheet["H2"].value, "\n\n0\n1.5")
            for coordinate in ("R2", "S2", "M3", "O3", "P3"):
                self.assertIsNone(sheet[coordinate].value)
        finally:
            workbook.close()

    async def test_long_application_text_fails_instead_of_silently_truncating(self):
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=[{
            "numero_refax": "0000001", "precio": 1, "stock": "DISPONIBLE",
            "aplicaciones": "x" * 32768,
        }]))
        with self.assertRaisesRegex(RefaxProductsError, "APLICACIONES.*32.767"):
            await service.download()

    async def test_download_renews_and_retries_an_unauthorized_token(self):
        service, connection_service = self.build_service()
        connection_service.get_valid_token.side_effect = [
            "expired-token",
            "renewed-token",
        ]
        service._request_products = AsyncMock(
            side_effect=[
                httpx.Response(401, json={"message": "Unauthorized"}),
                httpx.Response(200, json=[{"numero_refax": "ABC", "precio": 100, "stock": "Disponible"}]),
            ]
        )

        result = await service.download()

        self.assertTrue(result.content)
        connection_service.connect.assert_awaited_once()
        self.assertEqual(
            [call.args[0] for call in service._request_products.await_args_list],
            ["expired-token", "renewed-token"],
        )

    async def test_download_reports_an_upstream_error_without_leaking_body(self):
        service, _connection_service = self.build_service()
        service._request_products = AsyncMock(
            return_value=httpx.Response(
                500,
                text="internal details that must not reach the browser",
            )
        )

        with self.assertRaisesRegex(RefaxProductsError, "HTTP 500"):
            await service.download()

    async def test_invalid_payloads_fail_instead_of_downloading_error_as_excel(self):
        payloads = [
            b'<html>Gateway error</html>',
            b'{"error":"REFAX failed"}',
            b'[]',
            b'[{"numero_refax":"ABC","stock":"5"}]',
            b'[{"numero_refax":"ABC","precio":null,"stock":"5"}]',
            b'[{"numero_refax":"ABC","precio":"NaN","stock":"5"}]',
        ]
        service, _ = self.build_service()
        for payload in payloads:
            with self.subTest(payload=payload):
                service._request_products = AsyncMock(
                    return_value=httpx.Response(200, content=payload)
                )
                with self.assertRaises(RefaxProductsError):
                    await service.download()

    async def test_persistent_unauthorized_response_refreshes_only_once(self):
        service, connection_service = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(401))
        with self.assertRaisesRegex(RefaxProductsError, "HTTP 401"):
            await service.download()
        connection_service.connect.assert_awaited_once()
        self.assertEqual(service._request_products.await_count, 2)

    async def test_transport_failure_retries_with_same_code_and_bearer(self):
        service, connection_service = self.build_service()
        requests = []

        async def handle(request):
            requests.append(request)
            if len(requests) == 1:
                raise httpx.ReadError("connection interrupted", request=request)
            return httpx.Response(200, json=[
                {"numero_refax": "00123", "precio": 100, "stock": "Disponible"}
            ])

        client_class = httpx.AsyncClient
        def client_factory(**kwargs):
            return client_class(transport=httpx.MockTransport(handle), **kwargs)

        with (
            patch("services.refax_products_service.httpx.AsyncClient", side_effect=client_factory),
            patch("services.refax_products_service.asyncio.sleep", new_callable=AsyncMock) as sleep,
        ):
            result = await service.download()

        self.assertEqual(result.content_type, XLSX_CONTENT_TYPE)
        self.assertEqual(len(requests), 2)
        for request in requests:
            self.assertEqual(request.method, "GET")
            self.assertEqual(request.url.path, "/api/Productos/Listado")
            self.assertEqual(request.url.params["codigo"], settings.refax_provider_code)
            self.assertEqual(request.headers["Authorization"], "Bearer private-token")
        sleep.assert_awaited_once_with(1.0)
        connection_service.connect.assert_not_awaited()

    async def test_timeouts_stop_after_three_attempts_without_leaking_secrets(self):
        service, _ = self.build_service()
        with (
            patch("services.refax_products_service.httpx.AsyncClient") as client,
            patch("services.refax_products_service.asyncio.sleep", new_callable=AsyncMock) as sleep,
        ):
            get = client.return_value.__aenter__.return_value.get
            get.side_effect = httpx.ReadTimeout("private-token in unsafe error message")
            with self.assertLogs("services.refax_products_service", level="WARNING") as logs:
                with self.assertRaises(RefaxProductsError) as raised:
                    await service.download()
        self.assertIn("tiempo de espera", str(raised.exception))
        self.assertNotIn("private-token", str(raised.exception))
        self.assertNotIn("private-token", " ".join(logs.output))
        self.assertIn("ReadTimeout", " ".join(logs.output))
        self.assertEqual(get.await_count, 3)
        self.assertEqual(sleep.await_count, 2)

    async def test_temporary_http_errors_retry_but_permanent_errors_do_not(self):
        service, _ = self.build_service()
        with (
            patch("services.refax_products_service.httpx.AsyncClient") as client,
            patch("services.refax_products_service.asyncio.sleep", new_callable=AsyncMock) as sleep,
        ):
            get = client.return_value.__aenter__.return_value.get
            get.side_effect = [
                httpx.Response(503),
                httpx.Response(429, headers={"Retry-After": "3"}),
                httpx.Response(200, json=[]),
            ]
            result = await service._request_products("private-token")
            self.assertEqual(result.status_code, 200)
            self.assertEqual([call.args[0] for call in sleep.await_args_list], [1.0, 3.0])
            get.reset_mock(side_effect=True)
            get.return_value = httpx.Response(400)
            result = await service._request_products("private-token")
            self.assertEqual(result.status_code, 400)
            get.assert_awaited_once()

    async def test_route_serves_xlsx_attachment_and_json_error_on_failure(self):
        service, _ = self.build_service()
        service._request_products = AsyncMock(return_value=httpx.Response(
            200, json=[{"numero_refax": "ABC", "precio": 100, "stock": "Disponible"}]
        ))
        app = FastAPI()
        app.include_router(router)
        with patch("routers.refax_router.refax_products_service", service):
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://test"
            ) as client:
                response = await client.get("/refax/products/download")
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers["content-type"], XLSX_CONTENT_TYPE)
                self.assertTrue(response.headers["content-disposition"].endswith('.xlsx"'))
                self.assertEqual(response.headers["cache-control"], "no-store")
                self.assertTrue(response.content.startswith(b"PK"))
                service._request_products = AsyncMock(return_value=httpx.Response(200, json={"error": "bad"}))
                response = await client.get("/refax/products/download")
                self.assertEqual(response.status_code, 502)
                self.assertIn("detail", response.json())
                self.assertNotIn("content-disposition", response.headers)


if __name__ == "__main__":
    unittest.main()
