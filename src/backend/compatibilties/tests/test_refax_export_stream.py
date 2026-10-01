import asyncio
import base64
import json
import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook

from routers.refax_router import router
from services.refax_export_stream import stream_refax_export
from services.refax_products_service import RefaxProductsService, RefaxProductsDownload, RefaxProductsError, XLSX_CONTENT_TYPE


class RefaxExportStreamTests(unittest.IsolatedAsyncioTestCase):
    async def test_reports_progress_before_result_and_cancels_on_disconnect(self):
        cancelled = asyncio.Event()

        async def download(on_progress):
            on_progress(5, "requesting", "Consultando REFAX")
            try:
                await asyncio.Event().wait()
            finally:
                cancelled.set()

        stream = stream_refax_export(SimpleNamespace(download=download))
        first = json.loads(await anext(stream))
        self.assertEqual(first["percentage"], 0)
        progress = json.loads(await asyncio.wait_for(anext(stream), timeout=1))
        self.assertEqual(progress["percentage"], 5)
        await stream.aclose()
        self.assertTrue(cancelled.is_set())

    async def test_stream_route_sends_row_progress_and_exact_excel(self):
        products = [{"numero_refax": str(i).zfill(5), "precio": 1500, "stock": "Disponible"} for i in range(100)]
        connection = SimpleNamespace(get_valid_token=AsyncMock(return_value="private-token"))
        service = RefaxProductsService(connection)
        service._request_products = AsyncMock(return_value=httpx.Response(200, json=products))
        app = FastAPI()
        app.include_router(router)
        with patch("routers.refax_router.refax_products_service", service):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
                response = await client.get("/refax/products/download?progress=true")
        self.assertEqual(response.headers["content-type"], "application/x-ndjson")
        self.assertEqual(response.headers["x-accel-buffering"], "no")
        events = [json.loads(line) for line in response.text.splitlines()]
        values = [event["percentage"] for event in events if event["type"] == "progress"]
        self.assertTrue(any(60 < value < 90 for value in values))
        self.assertEqual(values, sorted(values))
        self.assertEqual(events[-1]["type"], "complete")
        metadata = next(e for e in events if e["type"] == "file")
        content = b"".join(base64.b64decode(e["data"]) for e in events if e["type"] == "chunk")
        self.assertEqual(len(content), metadata["size"])
        self.assertEqual(metadata["content_type"], XLSX_CONTENT_TYPE)
        self.assertTrue(metadata["filename"].endswith(".xlsx"))
        workbook = load_workbook(BytesIO(content))
        try:
            self.assertEqual(workbook.active.max_row, 101)
            self.assertEqual(workbook.active["A2"].value, "00000")
            self.assertEqual(workbook.active["C2"].value, "Disponible")
        finally:
            workbook.close()
        self.assertNotIn("private-token", response.text)
        service._request_products.assert_awaited_once()

    async def test_errors_do_not_produce_file_or_completion(self):
        for error in (RefaxProductsError("REFAX no disponible"), RuntimeError("private-secret")):
            service = SimpleNamespace(download=AsyncMock(side_effect=error))
            events = [json.loads(line) async for line in stream_refax_export(service)]
            self.assertEqual(events[-1]["type"], "error")
            self.assertFalse(any(e["type"] in ("file", "complete") for e in events))
            self.assertNotIn("private-secret", json.dumps(events))

    async def test_chunked_upstream_reports_bytes_and_measured_rows(self):
        content = json.dumps([{"numero_refax": "001", "precio": 20, "stock": "Disponible", "extra": "x" * 150000}]).encode()

        class Body(httpx.AsyncByteStream):
            async def __aiter__(self):
                for offset in range(0, len(content), 32768):
                    yield content[offset:offset + 32768]

        for known_size in (False, True):
            with self.subTest(known_size=known_size):
                def handle(request):
                    self.assertEqual(request.headers["authorization"], "Bearer private-token")
                    headers = {"content-length": str(len(content))} if known_size else {}
                    return httpx.Response(200, headers=headers, stream=Body())

                real_client = httpx.AsyncClient
                def client_factory(**kwargs):
                    return real_client(transport=httpx.MockTransport(handle), **kwargs)

                service = RefaxProductsService(SimpleNamespace(get_valid_token=AsyncMock(return_value="private-token")))
                progress = []
                with patch("services.refax_products_service.httpx.AsyncClient", side_effect=client_factory):
                    result = await service.download(on_progress=lambda *event: progress.append(event))
                received = [event for event in progress if event[1] == "receiving"]
                self.assertGreater(len(received), 2)
                self.assertTrue(any("MB" in event[2] for event in received))
                if known_size:
                    self.assertTrue(any(10 < event[0] <= 55 for event in received))
                else:
                    self.assertTrue(all(event[0] == 10 for event in received))
                self.assertEqual(progress[-1][0], 95)
                self.assertTrue(result.content.startswith(b"PK"))
