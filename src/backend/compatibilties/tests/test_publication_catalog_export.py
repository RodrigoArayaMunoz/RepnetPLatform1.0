import os
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

import httpx
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from config import settings
from routers import publications_router
from services.job_store import JobStore
from services.ml_client import ml_client
from services.publication_export_service import PublicationExportService
from services.publication_export_store import PublicationExportStore
from services.supabase_publications_store import (
    SupabasePublicationsStore,
    supabase_publications_store,
)
from tasks.publication_export_tasks import _export_publications_task


class CatalogDatabaseTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.enterContext(patch.object(settings, "supabase_url", "https://database.example.test"))
        self.store = SupabasePublicationsStore()
        self.store._headers = Mock(return_value={})
        self.requests = []

    def _client(self, responses):
        outcomes = iter(responses)

        def respond(request):
            self.requests.append(request)
            return next(outcomes)

        return httpx.AsyncClient(transport=httpx.MockTransport(respond))

    async def test_count_includes_all_dates_sellers_and_statuses(self):
        client = self._client([httpx.Response(206, headers={"content-range": "0-0/1005"})])
        with patch("services.supabase_publications_store.httpx.AsyncClient", return_value=client):
            count = await self.store.count_all()
        self.assertEqual(count, 1005)
        self.assertEqual(dict(self.requests[0].url.params), {"select": "mlc"})

    async def test_smaller_server_pages_do_not_truncate_the_catalog(self):
        client = self._client([
            httpx.Response(206, json=[{"mlc": "MLC1"}, {"mlc": "MLC2"}],
                           headers={"content-range": "0-1/5"}),
            httpx.Response(206, json=[{"mlc": "MLC3"}, {"mlc": "MLC4"}],
                           headers={"content-range": "2-3/*"}),
            httpx.Response(206, json=[{"mlc": "MLC5"}],
                           headers={"content-range": "4-4/*"}),
        ])
        with patch("services.supabase_publications_store.httpx.AsyncClient", return_value=client):
            rows = await self.store.list_all()
        self.assertEqual([row["mlc"] for row in rows], ["MLC1", "MLC2", "MLC3", "MLC4", "MLC5"])
        self.assertEqual([r.headers["range"] for r in self.requests], ["0-999", "2-1001", "4-1003"])
        for request in self.requests:
            self.assertEqual(set(request.url.params), {"select", "order"})
            self.assertEqual(request.url.params["order"], "seller_id.asc,mlc.asc")

    async def test_unknown_total_can_finish_with_an_out_of_range_response(self):
        client = self._client([
            httpx.Response(206, json=[{"mlc": "MLC1"}], headers={"content-range": "0-0/*"}),
            httpx.Response(416, json={"code": "PGRST103"}, headers={"content-range": "*/1"}),
        ])
        with patch("services.supabase_publications_store.httpx.AsyncClient", return_value=client):
            rows = await self.store.list_all()
        self.assertEqual(rows, [{"mlc": "MLC1"}])

    async def test_date_export_keeps_its_date_and_seller_filters(self):
        client = self._client([
            httpx.Response(200, json=[{"mlc": "MLC1"}], headers={"content-range": "0-0/1"}),
        ])
        with patch("services.supabase_publications_store.httpx.AsyncClient", return_value=client):
            await self.store.list_by_creation_date("2026-10-06", seller_id="99")
        params = self.requests[0].url.params
        self.assertEqual(params["fecha_creacion"], "eq.2026-10-06")
        self.assertEqual(params["seller_id"], "eq.99")

    async def test_bad_or_failed_pages_never_return_a_partial_catalog(self):
        for response in (httpx.Response(200, json=["invalid-row"]), httpx.Response(503)):
            with self.subTest(status=response.status_code):
                client = self._client([
                    httpx.Response(206, json=[{"mlc": "MLC1"}], headers={"content-range": "0-0/2"}),
                    response,
                ])
                with (
                    patch("services.supabase_publications_store.httpx.AsyncClient", return_value=client),
                    self.assertRaises(RuntimeError),
                ):
                    await self.store.list_all()


class CatalogExcelTests(unittest.IsolatedAsyncioTestCase):
    async def test_all_1005_database_rows_reach_the_excel_with_the_existing_columns(self):
        publications = [
            {
                "mlc": f"MLC{index:05}",
                "sku": f"SKU-{index}",
                "part_number": f"00{index}-A",
                "titulo": f"Publicación {index}",
                "status": "paused" if index % 2 else "active",
                "has_compatibilities": bool(index % 2),
                "seller_id": 99 if index % 2 else 100,
                "fecha_creacion": "2026-10-06" if index % 2 else None,
            }
            for index in range(1005)
        ]
        requests = []

        def respond(request):
            requests.append(request)
            self.assertEqual(set(request.url.params), {"select", "order"})
            start, end = map(int, request.headers["range"].split("-"))
            batch = publications[start:end + 1]
            return httpx.Response(
                206, json=batch,
                headers={"content-range": f"{start}-{start + len(batch) - 1}/1005"},
            )

        client = httpx.AsyncClient(transport=httpx.MockTransport(respond))
        with tempfile.TemporaryDirectory() as temp_dir:
            with (
                patch.object(settings, "supabase_url", "https://database.example.test"),
                patch.object(settings, "upload_dir", temp_dir),
                patch.object(supabase_publications_store, "_headers", return_value={}),
                patch("services.supabase_publications_store.httpx.AsyncClient", return_value=client),
                patch.object(supabase_publications_store, "list_by_creation_date", new=AsyncMock()) as date_query,
                patch.object(ml_client, "request", new=AsyncMock()) as ml_request,
                patch.object(JobStore, "update") as update,
            ):
                summary = await PublicationExportService().export(
                    job_id="catalog-job", user_id=None, creation_date=None,
                )
            workbook = load_workbook(os.path.join(temp_dir, "catalog-job_catalogo_completo_emilia.xlsx"))
            try:
                sheet = workbook["Publicaciones"]
                rows = list(sheet.iter_rows(values_only=True))
                self.assertEqual(sheet.auto_filter.ref, "A1:F1006")
                self.assertEqual(sheet.freeze_panes, "A2")
            finally:
                workbook.close()
        self.assertEqual(rows[0], ("MLC", "SKU", "NUMERO_PIEZA", "TITULO", "ESTADO", "¿POSEE COMPATIBILIDADES?"))
        self.assertEqual(len(rows), 1006)
        self.assertEqual([row[0] for row in rows[1:]], [p["mlc"] for p in publications])
        self.assertEqual(rows[2], ("MLC00001", "SKU-1", "001-A", "Publicación 1", "pausada/inactiva", "Sí"))
        self.assertEqual(summary["total_rows"], 1005)
        self.assertEqual(summary["api_items_queried"], 0)
        self.assertEqual(update.call_args.kwargs["output_filename"], "catalogo_completo_emilia.xlsx")
        self.assertEqual(len(requests), 2)
        date_query.assert_not_awaited()
        ml_request.assert_not_awaited()

    async def test_empty_catalog_does_not_create_an_excel(self):
        with (
            patch.object(supabase_publications_store, "list_all", new=AsyncMock(return_value=[])),
            patch.object(JobStore, "update"),
            self.assertRaisesRegex(ValueError, "catalogo completo"),
        ):
            await PublicationExportService().export(job_id="empty", user_id=None, creation_date=None)

    async def test_worker_runs_the_catalog_with_no_date_or_seller_filter(self):
        with (
            patch.object(JobStore, "get", return_value={}),
            patch.object(JobStore, "update"),
            patch.object(publications_router.publication_export_store, "renew_run_lock", return_value=True),
            patch("tasks.publication_export_tasks.publication_export_service.export", new=AsyncMock()) as export,
        ):
            await _export_publications_task(
                job_id="catalog-job", user_id=None, creation_date=None, lock_owner="task-id",
            )
        self.assertIsNone(export.call_args.kwargs["creation_date"])
        self.assertIsNone(export.call_args.kwargs["user_id"])


class CatalogEndpointTests(unittest.TestCase):
    def setUp(self):
        app = FastAPI()
        app.include_router(publications_router.router)

        @app.middleware("http")
        async def authenticate(request, call_next):
            request.state.supabase_user = {"id": "user-a"}
            return await call_next(request)

        self.client = TestClient(app)
        self.enterContext(patch.object(settings, "backend_auth_enabled", True))

    def test_missing_date_is_rejected_unless_catalog_is_explicit(self):
        for payload in ({}, {"refresh": True}, {"export_scope": "invalid"},
                        {"export_scope": "catalog", "publication_date": "2026-10-06"}):
            with self.subTest(payload=payload):
                self.assertEqual(self.client.post("/publications/export", json=payload).status_code, 422)

    def test_empty_catalog_returns_a_clear_error(self):
        with patch.object(supabase_publications_store, "count_all", new=AsyncMock(return_value=0)):
            response = self.client.post("/publications/export", json={"export_scope": "catalog"})
        self.assertEqual(response.status_code, 404)
        self.assertIn("catalogo completo", response.json()["detail"])

    def test_catalog_is_queued_separately_without_requiring_a_live_ml_connection(self):
        jobs = {}

        def create(filename):
            jobs["catalog-job"] = {"id": "catalog-job", "filename": filename}
            return jobs["catalog-job"]

        def update(job_id, **values):
            jobs[job_id].update(values)

        with (
            patch.object(supabase_publications_store, "count_all", new=AsyncMock(return_value=1005)) as count_all,
            patch.object(supabase_publications_store, "count_by_creation_date", new=AsyncMock()) as count_date,
            patch.object(publications_router, "_get_connected_ml_user_id", new=AsyncMock()) as ml_connection,
            patch.object(publications_router, "_existing_export_job", return_value=None),
            patch.object(publications_router.publication_export_store, "claim_reference", return_value=True) as claim,
            patch.object(JobStore, "create", side_effect=create),
            patch.object(JobStore, "get", side_effect=lambda job_id: jobs.get(job_id)),
            patch.object(JobStore, "update", side_effect=update),
            patch.object(publications_router.export_publications_task, "delay",
                         return_value=SimpleNamespace(id="task-id")) as delay,
        ):
            response = self.client.post("/publications/export", json={"export_scope": "catalog", "refresh": True})
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(result["export_scope"], "catalog")
        self.assertIsNone(result["publication_date"])
        self.assertEqual(result["filename"], "catalogo_completo_emilia.xlsx")
        self.assertEqual(jobs["catalog-job"]["requested_by_user_id"], "user-a")
        delay.assert_called_once_with("catalog-job", None, None)
        claim.assert_called_once_with(
            requested_by_user_id="user-a", seller_id=None, creation_date=None, job_id="catalog-job",
        )
        count_all.assert_awaited_once_with()
        count_date.assert_not_awaited()
        ml_connection.assert_not_awaited()

    def test_catalog_and_date_references_cannot_collide(self):
        catalog = PublicationExportStore._reference_key("user-a", None, None)
        dated = PublicationExportStore._reference_key("user-a", "99", "2026-10-06")
        other_user = PublicationExportStore._reference_key("user-b", None, None)
        self.assertEqual(len({catalog, dated, other_user}), 3)

    def test_interrupted_catalog_can_be_requeued(self):
        job = {
            "id": "catalog-job", "status": "processing", "heartbeat_at": 100,
            "export_scope": "catalog", "publication_date": None, "ml_user_id": None,
        }
        with (
            patch("routers.publications_router.time.time", return_value=500),
            patch.object(publications_router.publication_export_store, "has_run_lock", return_value=False),
            patch.object(publications_router.publication_export_store, "try_acquire_recovery_guard", return_value=True),
            patch.object(publications_router.export_publications_task, "delay",
                         return_value=SimpleNamespace(id="recovery-task")) as delay,
            patch.object(JobStore, "update"),
            patch.object(JobStore, "get", return_value={**job, "status": "queued"}),
        ):
            result = publications_router._recover_stale_export_if_needed(job)
        self.assertEqual(result["status"], "queued")
        delay.assert_called_once_with("catalog-job", None, None)


if __name__ == "__main__":
    unittest.main()
